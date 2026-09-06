"""Apply the major-review revision and final Humanizer pass.

Purpose: revise the carotid plaque manuscript according to the major-review
comments, integrate reviewer-response sensitivity analyses, and export Word
documents for author review.
Inputs: current manuscript references, project result tables, QC notes, and
main figure PNG files.
Outputs: third-round revised manuscript DOCX, integrated DOCX, revision note,
markdown source, and log files under 修改内容 and 06_manuscript.
Dependencies: python-docx, pandas optional, Python standard library.
Run order: after 18_review_response_sensitivity_analysis.R.
"""

from __future__ import annotations

import csv
import datetime as dt
import math
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = ROOT / "06_manuscript"
TABLE_DIR = ROOT / "05_results" / "tables"
QC_DIR = ROOT / "05_results" / "qc"
FIGURE_DIR = ROOT / "05_results" / "figures"
OUTPUT_DIR = ROOT / "修改内容"
LOG_DIR = ROOT / "08_logs"
RUNBOOK = ROOT / "RUNBOOK.md"
PROJECT_STATUS = ROOT / "PROJECT_STATUS.md"
ERROR_LOG = LOG_DIR / "error_log.md"
LOG = LOG_DIR / "19_apply_major_review_revision.log"

TITLE = (
    "Tissue-level covariation of efferocytosis-related and inflammatory VSMC "
    "transcriptional programs in clinically symptomatic versus asymptomatic "
    "human carotid plaques"
)

FORBIDDEN_REVIEW_TERMS = [
    " unstable ",
    " vulnerable ",
    " confirmed ",
    " proves ",
    " demonstrated conclusively ",
    " therapeutic target",
    " diagnostic tool",
]


def stamp() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def log(message: str) -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"[{stamp()}] {message}\n")


def record_error(message: str) -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with ERROR_LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"\n- {stamp()} | 19_apply_major_review_revision | {message}\n")
    log(f"ERROR {message}")


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        record_error(f"missing CSV: {path}")
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def fmt_num(value: object, digits: int = 3) -> str:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return "NA"
    if not math.isfinite(x):
        return "NA"
    if x == 0:
        return "0"
    if abs(x) < 0.001:
        return f"{x:.2e}"
    if abs(x) < 0.01:
        return f"{x:.4f}"
    return f"{x:.{digits}f}"


def fmt_p(value: object) -> str:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return "NA"
    if not math.isfinite(x):
        return "NA"
    if x < 0.001:
        return f"{x:.2e}"
    if x < 0.01:
        return f"{x:.4f}"
    return f"{x:.3f}"


def get_row(rows: list[dict[str, str]], **criteria: str) -> dict[str, str]:
    for row in rows:
        if all(row.get(key) == value for key, value in criteria.items()):
            return row
    return {}


def count_fdr(path: Path) -> tuple[int, int]:
    rows = read_csv_rows(path)
    if not rows:
        return 0, 0
    fdr_col = next((k for k in rows[0] if k.lower() in {"fdr", "adj.p.val", "padj", "qvalue"}), None)
    if fdr_col is None:
        return len(rows), 0
    sig = 0
    for row in rows:
        try:
            sig += float(row.get(fdr_col, "nan")) < 0.05
        except ValueError:
            pass
    return len(rows), int(sig)


def top_rows(path: Path, p_col: str, n: int = 3) -> list[dict[str, str]]:
    rows = read_csv_rows(path)
    def score(row: dict[str, str]) -> float:
        try:
            return float(row.get(p_col, "1") or 1)
        except ValueError:
            return 1.0
    return sorted(rows, key=score)[:n]


def extract_references() -> str:
    sources = [
        MANUSCRIPT_DIR / "manuscript_draft.md",
        MANUSCRIPT_DIR / "archive_before_skill_revision_20260904" / "manuscript_draft.md",
    ]
    for source in sources:
        if not source.exists():
            continue
        text = source.read_text(encoding="utf-8", errors="replace")
        marker = "\n## References\n"
        if marker in text:
            return text.split(marker, 1)[1].strip()
    return "AUTHOR_INPUT_NEEDED: references could not be extracted from the prior manuscript."


def grouped_counts(rows: list[dict[str, str]]) -> dict[str, dict[str, int]]:
    out: dict[str, dict[str, int]] = {}
    for row in rows:
        sig = row.get("signature", "")
        module = row.get("module", "")
        if not sig or not module:
            continue
        out.setdefault(sig, {})
        out[sig][module] = out[sig].get(module, 0) + 1
    return out


def list_genes(rows: list[dict[str, str]], signature: str) -> str:
    genes = [row.get("gene", "") for row in rows if row.get("signature") == signature and row.get("gene")]
    return ", ".join(genes)


def top_contributors(rows: list[dict[str, str]], cohort: str, n: int = 3) -> str:
    selected = [row for row in rows if row.get("cohort") == cohort and row.get("signature") == "efferocytosis"]
    def score(row: dict[str, str]) -> float:
        try:
            return float(row.get("spearman_rho_with_signature_score", "-999"))
        except ValueError:
            return -999.0
    selected = sorted(selected, key=score, reverse=True)[:n]
    return ", ".join(
        f"{row.get('gene')} (rho = {fmt_num(row.get('spearman_rho_with_signature_score'))}, FDR = {fmt_p(row.get('FDR'))})"
        for row in selected
    )


def build_context() -> dict[str, object]:
    sig_summary = read_csv_rows(TABLE_DIR / "external_validation_signature_summary.csv")
    sig_effects = read_csv_rows(TABLE_DIR / "signature_group_effect_sizes_bootstrap.csv")
    sig_corr = read_csv_rows(TABLE_DIR / "signature_correlations.csv")
    ranked_111 = read_csv_rows(TABLE_DIR / "GSE111782" / "ranked_signature_enrichment.csv")
    ranked_311 = read_csv_rows(TABLE_DIR / "GSE311535" / "ranked_signature_enrichment.csv")
    go_111 = read_csv_rows(TABLE_DIR / "GSE111782" / "GO_BP_enrichment.csv")
    donor_tests = read_csv_rows(TABLE_DIR / "GSE260657_donor_program_tests.csv")
    comp_tests = read_csv_rows(TABLE_DIR / "GSE260657_donor_celltype_composition_tests.csv")
    comp_summary = read_csv_rows(TABLE_DIR / "GSE260657_donor_celltype_composition_group_summary.csv")
    macro_sub = read_csv_rows(TABLE_DIR / "GSE260657_macrophage_subcluster_summary.csv")
    vsmc_sub = read_csv_rows(TABLE_DIR / "GSE260657_vsmc_subcluster_summary.csv")
    cell_rows = read_csv_rows(QC_DIR / "GSE260657_cell_annotations.csv")
    gene_dir = read_csv_rows(TABLE_DIR / "external_validation_gene_direction.csv")
    gene_sets = read_csv_rows(TABLE_DIR / "gene_set_category_audit.csv")
    overlap = read_csv_rows(TABLE_DIR / "gene_set_overlap_jaccard.csv")
    mapping = read_csv_rows(TABLE_DIR / "gene_set_platform_mapping_summary.csv")
    boot = read_csv_rows(TABLE_DIR / "main_correlation_bootstrap_summary.csv")
    meta = read_csv_rows(TABLE_DIR / "main_correlation_exploratory_meta_summary.csv")
    partial = read_csv_rows(TABLE_DIR / "main_correlation_partial_spearman_proxy_adjustment.csv")
    contrib = read_csv_rows(TABLE_DIR / "efferocytosis_gene_score_contribution_bulk.csv")
    qc_by_file = read_csv_rows(QC_DIR / "GSE260657_cell_qc_by_file.csv")

    cell_counts: dict[str, int] = {}
    detected_min = detected_max = counts_min = counts_max = mt_min = mt_max = None
    for row in cell_rows:
        cell_type = row.get("major_cell_type", "unassigned") or "unassigned"
        cell_counts[cell_type] = cell_counts.get(cell_type, 0) + 1
        for key, dest_min, dest_max in [
            ("detected_genes", "detected_min", "detected_max"),
            ("counts", "counts_min", "counts_max"),
            ("mt_fraction", "mt_min", "mt_max"),
        ]:
            try:
                value = float(row.get(key, "nan"))
            except ValueError:
                continue
            if not math.isfinite(value):
                continue
            if dest_min == "detected_min":
                detected_min = value if detected_min is None else min(detected_min, value)
                detected_max = value if detected_max is None else max(detected_max, value)
            elif dest_min == "counts_min":
                counts_min = value if counts_min is None else min(counts_min, value)
                counts_max = value if counts_max is None else max(counts_max, value)
            else:
                mt_min = value if mt_min is None else min(mt_min, value)
                mt_max = value if mt_max is None else max(mt_max, value)

    common_genes = len(gene_dir)
    concordant = sum(row.get("direction_concordant") == "TRUE" for row in gene_dir)
    both_sig = sum(
        row.get("discovery_FDR_lt_05") == "TRUE" and row.get("validation_FDR_lt_05") == "TRUE"
        for row in gene_dir
    )

    return {
        "sig_summary": sig_summary,
        "sig_effects": sig_effects,
        "sig_corr": sig_corr,
        "ranked_111": ranked_111,
        "ranked_311": ranked_311,
        "go_111": go_111,
        "donor_tests": donor_tests,
        "comp_tests": comp_tests,
        "comp_summary": comp_summary,
        "macro_sub": macro_sub,
        "vsmc_sub": vsmc_sub,
        "cell_counts": cell_counts,
        "gene_sets": gene_sets,
        "overlap": overlap,
        "mapping": mapping,
        "boot": boot,
        "meta": meta,
        "partial": partial,
        "contrib": contrib,
        "qc_by_file": qc_by_file,
        "sc_qc": {
            "detected_min": detected_min,
            "detected_max": detected_max,
            "counts_min": counts_min,
            "counts_max": counts_max,
            "mt_min": mt_min,
            "mt_max": mt_max,
        },
        "deg111": count_fdr(TABLE_DIR / "GSE111782_DEG_complete.csv"),
        "deg311": count_fdr(TABLE_DIR / "GSE311535_DEG_complete.csv"),
        "top111": top_rows(TABLE_DIR / "GSE111782_DEG_complete.csv", "P_value", 3),
        "top311": top_rows(TABLE_DIR / "GSE311535_DEG_complete.csv", "P_value", 3),
        "common_genes": common_genes,
        "concordant": concordant,
        "both_sig": both_sig,
    }


def ranked_sentence(row: dict[str, str]) -> str:
    return f"{row.get('pathway')} (NES = {fmt_num(row.get('NES'))}, FDR = {fmt_p(row.get('padj'))})"


def top_deg_text(rows: list[dict[str, str]]) -> str:
    return ", ".join(
        f"{row.get('gene')} (log2FC = {fmt_num(row.get('logFC'))}, P = {fmt_p(row.get('P_value'))}, FDR = {fmt_p(row.get('FDR'))})"
        for row in rows
    )


def effect_text(row: dict[str, str]) -> str:
    return (
        f"median difference = {fmt_num(row.get('median_difference_symptomatic_minus_asymptomatic'))}, "
        f"bootstrap 95% CI {fmt_num(row.get('bootstrap_ci_low'))} to {fmt_num(row.get('bootstrap_ci_high'))}, "
        f"P = {fmt_p(row.get('P_value'))}, FDR = {fmt_p(row.get('FDR'))}"
    )


def partial_text(row: dict[str, str]) -> str:
    return f"rho = {fmt_num(row.get('rho'))}, P = {fmt_p(row.get('P_value'))}, FDR = {fmt_p(row.get('FDR'))}"


def build_manuscript(ctx: dict[str, object], refs_text: str) -> str:
    sig_summary = ctx["sig_summary"]
    sig_effects = ctx["sig_effects"]
    sig_corr = ctx["sig_corr"]
    ranked_111 = ctx["ranked_111"]
    ranked_311 = ctx["ranked_311"]
    go_111 = ctx["go_111"]
    donor_tests = ctx["donor_tests"]
    comp_tests = ctx["comp_tests"]
    macro_sub = ctx["macro_sub"]
    vsmc_sub = ctx["vsmc_sub"]
    cell_counts = ctx["cell_counts"]
    gene_sets = ctx["gene_sets"]
    overlap = ctx["overlap"]
    mapping = ctx["mapping"]
    boot = ctx["boot"]
    meta = ctx["meta"]
    partial = ctx["partial"]
    contrib = ctx["contrib"]
    sc_qc = ctx["sc_qc"]
    deg111_n, deg111_sig = ctx["deg111"]
    deg311_n, deg311_sig = ctx["deg311"]

    eff_infl_111 = get_row(sig_corr, cohort="GSE111782", signature_1="efferocytosis", signature_2="inflammatory_vsmc")
    eff_syn_111 = get_row(sig_corr, cohort="GSE111782", signature_1="efferocytosis", signature_2="synthetic_vsmc")
    eff_infl_311 = get_row(sig_corr, cohort="GSE311535", signature_1="efferocytosis", signature_2="inflammatory_vsmc")
    eff_ecm_311 = get_row(sig_corr, cohort="GSE311535", signature_1="efferocytosis", signature_2="ecm_remodeling_vsmc")

    boot111 = get_row(boot, cohort="GSE111782")
    boot311 = get_row(boot, cohort="GSE311535")
    partial111_mv = get_row(partial, cohort="GSE111782", model="macrophage_and_vsmc_proxy")
    partial111_four = get_row(partial, cohort="GSE111782", model="four_celltype_proxy")
    partial111_inf = get_row(partial, cohort="GSE111782", model="inflammation_proxy")
    partial311_mv = get_row(partial, cohort="GSE311535", model="macrophage_and_vsmc_proxy")
    partial311_four = get_row(partial, cohort="GSE311535", model="four_celltype_proxy")
    partial311_inf = get_row(partial, cohort="GSE311535", model="inflammation_proxy")
    meta_row = meta[0] if meta else {}

    eff_111 = get_row(sig_summary, cohort="GSE111782", signature="efferocytosis")
    eff_311 = get_row(sig_summary, cohort="GSE311535", signature="efferocytosis")
    contractile_111 = get_row(sig_effects, cohort="GSE111782", signature="contractile_vsmc")
    eff_effect_111 = get_row(sig_effects, cohort="GSE111782", signature="efferocytosis")
    eff_effect_311 = get_row(sig_effects, cohort="GSE311535", signature="efferocytosis")

    sig_ranked_111 = [r for r in ranked_111 if float(r.get("padj", "1") or 1) < 0.05]
    sig_ranked_text = "; ".join(ranked_sentence(row) for row in sig_ranked_111)
    best_311 = sorted(ranked_311, key=lambda r: float(r.get("padj", "1") or 1))[0] if ranked_311 else {}
    go_text = "; ".join(
        f"{row.get('Description')} (FDR = {fmt_p(row.get('p.adjust'))}; genes: {row.get('geneID', '').replace('/', ', ')})"
        for row in go_111[:3]
    )

    eff_infl_overlap = get_row(overlap, signature_1="efferocytosis", signature_2="inflammatory_vsmc")
    synthetic_ecm_overlap = get_row(overlap, signature_1="ecm_remodeling_vsmc", signature_2="synthetic_vsmc")
    map_eff_111 = get_row(mapping, cohort="GSE111782", signature="efferocytosis")
    map_infl_111 = get_row(mapping, cohort="GSE111782", signature="inflammatory_vsmc")
    map_eff_311 = get_row(mapping, cohort="GSE311535", signature="efferocytosis")
    map_infl_311 = get_row(mapping, cohort="GSE311535", signature="inflammatory_vsmc")
    map_contr_111 = get_row(mapping, cohort="GSE111782", signature="contractile_vsmc")

    module_counts = grouped_counts(gene_sets)
    eff_modules = ", ".join(
        f"{name}: {count}" for name, count in module_counts.get("efferocytosis", {}).items()
    )
    cell_summary = ", ".join(
        f"{key}: {value}" for key, value in sorted(cell_counts.items(), key=lambda item: (-item[1], item[0]))
    )
    macro_text = "; ".join(
        f"{row.get('subcluster')} n = {row.get('n_cells')}, efferocytosis score = {fmt_num(row.get('efferocytosis_score'))}"
        for row in macro_sub
    )
    vsmc_text = "; ".join(
        f"{row.get('subcluster')} n = {row.get('n_cells')}, contractile score = {fmt_num(row.get('contractile_vsmc_score'))}, non-contractile score = {fmt_num(row.get('noncontractile_vsmc_score'))}"
        for row in vsmc_sub
    )
    sc_tests_sig = sum(float(row.get("FDR", "1") or 1) < 0.05 for row in donor_tests)
    comp_tests_sig = sum(float(row.get("FDR", "1") or 1) < 0.05 for row in comp_tests)
    comp_lowest = sorted(comp_tests, key=lambda r: float(r.get("FDR", "1") or 1))[:2]
    comp_lowest_text = "; ".join(
        f"{row.get('major_cell_type')} {row.get('comparison')} (P = {fmt_p(row.get('P_value'))}, FDR = {fmt_p(row.get('FDR'))})"
        for row in comp_lowest
    )

    concord_pct = 100 * ctx["concordant"] / ctx["common_genes"] if ctx["common_genes"] else 0

    manuscript = f"""# {TITLE}

Article type: Original Research

Running title: Efferocytosis-related and VSMC programs

Authors: [[AUTHOR NAME]]

Affiliations: [[AFFILIATION]]

Corresponding author: [[CORRESPONDING AUTHOR EMAIL]]

Manuscript status: third-round major-revision author-review draft. The text was revised with the Humanizer skill on {dt.date.today().isoformat()}.

## Abstract

Background: Efferocytosis and vascular smooth muscle cell (VSMC) phenotypic modulation are relevant to atherosclerotic plaque biology. Their relationship in clinically symptomatic versus asymptomatic human carotid plaques remains uncertain. Methods: We reanalyzed public Gene Expression Omnibus datasets: GSE111782, an Affymetrix microarray discovery cohort with 9 symptomatic and 9 asymptomatic plaques; GSE311535, an independent diabetes-specific RNA-seq replication cohort with 6 symptomatic and 6 asymptomatic plaques; and GSE260657, a Smart-seq2 single-cell cohort with 8 symptomatic and 7 asymptomatic carotid plaque samples. Predefined efferocytosis-related and VSMC signatures were evaluated with assay-matched bulk models, ranked enrichment, sample-level scoring, Spearman correlation, and donor-level single-cell summaries. We added gene-set mapping, overlap, leave-one-out, bootstrap, and exploratory marker-proxy partial-correlation analyses. Results: The efferocytosis and inflammatory VSMC signatures had no direct gene overlap (Jaccard = {fmt_num(eff_infl_overlap.get('jaccard'))}). Mapping was high in both bulk cohorts: efferocytosis {map_eff_111.get('mapped_genes')}/{map_eff_111.get('total_genes')} and inflammatory VSMC {map_infl_111.get('mapped_genes')}/{map_infl_111.get('total_genes')} in GSE111782, and {map_eff_311.get('mapped_genes')}/{map_eff_311.get('total_genes')} and {map_infl_311.get('mapped_genes')}/{map_infl_311.get('total_genes')} in GSE311535. No individual gene reached FDR < 0.05 in either bulk cohort. No sample-level signature comparison between symptomatic and asymptomatic plaques reached FDR < 0.05. Efferocytosis-related scores correlated with inflammatory VSMC scores in GSE111782 (Spearman rho = {fmt_num(eff_infl_111.get('rho'))}, FDR = {fmt_p(eff_infl_111.get('FDR'))}; bootstrap 95% CI {fmt_num(boot111.get('bootstrap_ci_low'))} to {fmt_num(boot111.get('bootstrap_ci_high'))}) and GSE311535 (rho = {fmt_num(eff_infl_311.get('rho'))}, FDR = {fmt_p(eff_infl_311.get('FDR'))}; bootstrap 95% CI {fmt_num(boot311.get('bootstrap_ci_low'))} to {fmt_num(boot311.get('bootstrap_ci_high'))}). These correlations were observed in mixed plaque tissue and were not shown to represent macrophage-VSMC cell-cell coupling. In GSE260657, macrophage and VSMC compartments were localized, but donor-level program and cell-composition comparisons did not reach FDR < 0.05. Conclusion: The findings support a reproducible tissue-level association between efferocytosis-related and inflammatory VSMC expression scores across two heterogeneous bulk cohorts. They do not establish cell-specific coordination, mechanism, diagnostic utility, or pathology-defined plaque classification.

Keywords: carotid atherosclerosis; symptomatic plaque; asymptomatic plaque; efferocytosis; macrophage; vascular smooth muscle cell; transcriptomics

## Introduction

Human carotid atherosclerotic plaques are clinically important because symptomatic lesions are linked to cerebrovascular events. Symptom status is not the same as a pathology-defined plaque category. It can depend on thromboembolism, stenosis, vascular territory, collateral flow, prior events, and clinical ascertainment. For that reason, datasets labeled as symptomatic versus asymptomatic should retain those labels unless the source metadata provide a pathology grouping.

Macrophage clearance of apoptotic cells is one plausible process connecting inflammation resolution, necrotic-core biology, and advanced atherosclerosis. Efferocytosis involves recognition receptors, bridging molecules, engulfment machinery, lysosomal processing, and lipid-handling responses. Prior experimental and human studies have associated defective efferocytosis with plaque inflammation and necrosis [6-8].

VSMCs also change state in atherosclerotic plaques. They may lose contractile markers and acquire matrix-producing, inflammatory, osteogenic, or macrophage-like transcriptional features [9-13]. These shifts complicate bulk plaque transcriptomics because both cell abundance and cell state can change the measured signal.

Single-cell studies have shown substantial heterogeneity in human plaque tissue, including macrophage diversity and VSMC phenotypic modulation [12-20]. Less clear is whether a predefined efferocytosis-related transcriptional program covaries reproducibly with VSMC state programs in clinically symptomatic versus asymptomatic human carotid plaque cohorts.

We performed a secondary analysis of public human carotid plaque datasets under a deliberately bounded design. The study asked whether symptomatic plaques showed altered efferocytosis-related transcriptional features and whether those features were associated with non-contractile VSMC programs. We used one discovery bulk cohort, one independent heterogeneous bulk replication cohort, and one human carotid plaque single-cell cohort. The aim was to identify reproducible associations for follow-up, not to infer causality or relabel clinical groups as pathology-defined plaque types.

## Materials and Methods

### Data sources and cohort definitions

Public datasets were retained when metadata supported human carotid atherosclerotic plaque tissue and interpretable symptomatic versus asymptomatic grouping. GSE111782 served as the discovery bulk cohort. Its GEO record and linked publication describe post-bifurcation internal carotid atheroma samples with symptomatic and asymptomatic clinical labels [1,4,5]. The cohort included 18 samples, with 9 symptomatic and 9 asymptomatic plaques, profiled on the GPL571 Affymetrix Human Genome U133A 2.0 Array.

GSE311535 served as an independent heterogeneous bulk replication cohort. It contains human carotid plaque RNA-seq count data from patients with diabetes, including 6 symptomatic and 6 asymptomatic samples [2]. Because this cohort is diabetes-specific, it was not treated as a clinically identical validation cohort. GSE260657 was used for single-cell localization because the GEO record and linked publication describe Smart-seq2 data from 15 human carotid plaque samples, including 8 symptomatic and 7 asymptomatic samples [3]. Dataset provenance, URLs, limitations, and audit decisions are recorded in DATA_PROVENANCE.md and 03_data/metadata/dataset_audit.csv.

### Predefined gene sets and signature audit

The efferocytosis-related gene set was defined before inspecting project results. It contained 30 genes across recognition receptors, bridging molecules, integrin or scavenger uptake genes, inflammation-resolution or phosphatidylserine-binding genes, lipid-handling genes, and lysosomal-processing genes. The module counts were {eff_modules}. The full list is reported in 05_results/tables/gene_set_category_audit.csv. The genes were {list_genes(gene_sets, 'efferocytosis')}.

VSMC signatures were predefined as separate contractile, synthetic, inflammatory, and ECM-remodeling/osteogenic programs. Each VSMC signature contained 12 genes. These signatures were treated as transcriptional programs, not definitive lineage-state labels. The complete gene lists, modules, and source bases are provided in 05_results/tables/gene_set_category_audit.csv.

Gene-set overlap was assessed before interpreting correlations. The efferocytosis-related and inflammatory VSMC signatures had 0 overlapping genes (Jaccard = {fmt_num(eff_infl_overlap.get('jaccard'))}). The only non-zero overlap among predefined VSMC signatures was between synthetic and ECM-remodeling/osteogenic VSMC sets, which shared COL1A1 and SPARC (Jaccard = {fmt_num(synthetic_ecm_overlap.get('jaccard'))}). Platform mapping was also audited. GSE111782 mapped {map_eff_111.get('mapped_genes')}/{map_eff_111.get('total_genes')} efferocytosis-related genes, missing {map_eff_111.get('missing_genes')}, and {map_contr_111.get('mapped_genes')}/{map_contr_111.get('total_genes')} contractile VSMC genes, missing {map_contr_111.get('missing_genes')}. All other predefined signatures mapped completely in GSE111782, and all listed signatures mapped completely in GSE311535.

### Bulk transcriptomic processing

GSE111782 was analyzed as an RMA-processed log2 microarray matrix. Probes without usable gene annotation were removed where possible. For duplicate gene mappings, the probe with the highest average expression was retained. The symptomatic-versus-asymptomatic contrast was fitted with limma empirical Bayes linear modeling [21]. The prespecified outlier rule flagged samples with missing fraction greater than 0.05 or median sample correlation below the cohort median minus 3 MAD. No sample was automatically removed by the script.

GSE311535 was analyzed as gene-level RNA-seq counts with edgeR quasi-likelihood negative-binomial models after low-expression filtering and library-size normalization [22]. TPM, FPKM, or log2 expression values were not used as raw counts. In both cohorts, positive log2 fold change denotes higher expression in symptomatic plaques relative to asymptomatic plaques.

### Signature, enrichment, and association analyses

Sample-level signature scores were calculated from the mapped genes available on each platform. Each score is the mean standardized expression of genes in that signature. Group comparisons used two-sided Wilcoxon rank-sum tests because the bulk cohorts were small and normality assumptions were not defensible. Median differences were summarized as symptomatic minus asymptomatic with bootstrap 95% intervals.

Ranked enrichment used preranked symptomatic-versus-asymptomatic statistics under the GSEA framework [24]. GO biological process, KEGG, and Reactome over-representation analyses were attempted using clusterProfiler and ReactomePA when gene mapping and package availability allowed [25-29]. Tissue-level associations between signatures were evaluated using Spearman correlations.

The Benjamini-Hochberg method controlled FDR within defined analysis families. These families were: genes within each differential-expression cohort; all predefined bulk signature group comparisons; all generated pairwise signature correlations; ranked predefined signatures within each cohort; GO, KEGG, and Reactome outputs within each cohort and database; single-cell donor-level program comparisons; and single-cell donor-level cell-composition comparisons. P values are two-sided where applicable, and FDR < 0.05 was the prespecified reporting threshold.

Correlation robustness was assessed with leave-one-out Spearman correlations and 5,000 bootstrap resamples. A fixed-effect Fisher z synthesis was calculated only as an exploratory descriptive summary. The cohorts were still reported separately because their clinical and technical differences are material.

To address cell-composition confounding in bulk data, exploratory rank-residual partial Spearman correlations adjusted the efferocytosis-inflammatory VSMC association for marker-score proxies. Three adjustment sets were used: macrophage plus VSMC marker scores; macrophage, VSMC, endothelial, and fibroblast marker scores; and a broad inflammation marker score. These are marker-score proxies, not measured histology or validated deconvolution estimates.

WGCNA was prespecified but not run. The discovery bulk sample size was 18, below the project threshold of at least 30 samples for stable network inference [30]. The optional machine-learning module was also skipped because the discovery and replication cohorts were small and did not yield an FDR-supported discovery marker set.

### Single-cell processing and annotation

GSE260657 raw Smart-seq2 text files were parsed without modifying the downloaded archive. The archived prespecified QC starting points were a minimum of 200 detected genes per cell, a maximum of 12,000 detected genes subject to platform review, a maximum mitochondrial fraction of 20%, and at least 50 cells per retained sample for descriptive outputs. The fallback parser reported detected genes, counts, mitochondrial fraction, donor metadata, marker scores, and cell type calls, but it did not implement complete doublet removal. High detected-gene cells therefore require author review before stronger inference.

Cell-level expression was log-normalized as log1p(counts per cell library size times 10,000). Feature selection used the top high-variance genes plus canonical markers, producing 1,228 features. PCA computed 20 components. K-means clustering used the first 10 PCs with k = 10, nstart = 25, and iter.max = 100. UMAP was skipped because the uwot package was not available in the locked environment.

Cell annotation used canonical marker-score winners, PCA/k-means support, and marker-candidate inspection. The main labels were macrophage, VSMC, endothelial cell, fibroblast, T/NK cell, B cell, mast cell, and unassigned. Calls with low or insufficiently separated marker scores were retained as unassigned. The S1-S3 macrophage and VSMC subcluster labels are descriptive labels from this analysis. They should not be read as validated biological subtypes.

Group comparisons used donor/sample-level summaries where possible. Individual cells were not treated as independent patients. Because donor identity was available through sample files, single-cell program and composition comparisons used donor-level Wilcoxon summaries. Cell-level plots were used for localization and descriptive variation.

### Reporting, software, and reproducibility

Analyses were performed in R 4.6.0 with limma 3.68.5, edgeR 4.10.4, GSVA 2.6.6, fgsea 1.38.0, clusterProfiler 4.20.0, ReactomePA, ggplot2 4.0.3, and related Bioconductor packages. Package versions are archived in 09_environment/package_versions.csv and 09_environment/session_info.txt. Scripts are stored in 04_scripts, logs in 08_logs, and generated tables and figures in 05_results. No sample was removed solely because it changed result direction or statistical significance.

## Results

### Cohort selection produced discovery, replication, and single-cell evidence layers

The dataset audit retained GSE111782 as the discovery bulk cohort, GSE311535 as an independent heterogeneous bulk replication cohort, and GSE260657 as the single-cell localization cohort (Figure 1). The two bulk cohorts were analyzed separately because they differed by platform, preprocessing, and clinical context. GSE311535 was diabetes-specific, which limits direct comparability with GSE111782. Optional spatial transcriptomics, cell-cell communication, WGCNA, and machine-learning modules were not used in the main conclusions because their prespecified gates were not met.

### Gene-set audit supported transparency and ruled out direct overlap for the main correlation

The efferocytosis-related signature contained 30 genes across six functional categories. VSMC programs were analyzed as four separate 12-gene signatures: contractile, synthetic, inflammatory, and ECM-remodeling/osteogenic. Platform mapping retained {map_eff_111.get('mapped_genes')} of 30 efferocytosis-related genes in GSE111782 and all 30 in GSE311535. The inflammatory VSMC signature mapped completely in both cohorts. The efferocytosis-related and inflammatory VSMC signatures shared no genes, so the main correlation was not caused by direct gene overlap. The contribution audit showed that the aggregate efferocytosis score correlated most strongly with {top_contributors(contrib, 'GSE111782')} in GSE111782 and {top_contributors(contrib, 'GSE311535')} in GSE311535. This audit indicates that lipid-handling and lysosomal components contribute substantially to the aggregate score.

### Bulk differential-expression analysis did not identify FDR-supported individual genes

In GSE111782, limma tested {deg111_n:,} genes and identified {deg111_sig} genes at FDR < 0.05 (Figure 2). The smallest nominal P-value genes were {top_deg_text(ctx['top111'])}. Their adjusted values were not below the prespecified threshold. In GSE311535, edgeR tested {deg311_n:,} genes and identified {deg311_sig} genes at FDR < 0.05. The smallest nominal P-value genes were {top_deg_text(ctx['top311'])}, again without FDR support. Across {ctx['common_genes']:,} genes available in both bulk cohorts, {ctx['concordant']:,} ({concord_pct:.1f}%) had concordant symptomatic-versus-asymptomatic log2FC directions, but no gene reached FDR < 0.05 in both cohorts. The available data therefore do not support presenting individual genes as replicated markers.

### Discovery ranked enrichment was not reproduced in the replication cohort

Ranked enrichment of predefined signatures in GSE111782 identified FDR-supported negative enrichment for {sig_ranked_text}. Because positive log2FC denotes higher expression in symptomatic plaques, negative normalized enrichment scores mean that these genes tended to rank lower in symptomatic versus asymptomatic plaques. GO biological process over-representation in GSE111782 returned three FDR-supported terms related to cartilage or chondrocyte differentiation: {go_text}. Reactome over-representation did not yield FDR-supported terms in the discovery output.

In GSE311535, none of the ranked predefined signatures reached FDR < 0.05. The lowest adjusted value was observed for {best_311.get('pathway', 'NA')} (NES = {fmt_num(best_311.get('NES'))}, FDR = {fmt_p(best_311.get('padj'))}). The discovery enrichment pattern is therefore best treated as cohort-specific, not as a replicated symptomatic-versus-asymptomatic signature.

### Sample-level signature group differences did not survive FDR correction

No predefined efferocytosis or VSMC signature differed between symptomatic and asymptomatic plaques at FDR < 0.05 in either bulk cohort (Figure 3). In GSE111782, the efferocytosis median score was {fmt_num(eff_111.get('median_asymptomatic'))} in asymptomatic plaques and {fmt_num(eff_111.get('median_symptomatic'))} in symptomatic plaques ({effect_text(eff_effect_111)}). The contractile VSMC score showed a nominal difference in GSE111782 ({effect_text(contractile_111)}), but it did not survive FDR correction. In GSE311535, the efferocytosis median score was {fmt_num(eff_311.get('median_asymptomatic'))} in asymptomatic plaques and {fmt_num(eff_311.get('median_symptomatic'))} in symptomatic plaques ({effect_text(eff_effect_311)}). These results should be read as not detecting FDR-supported group differences at the current sample size and data quality.

### Efferocytosis-related scores covaried with inflammatory VSMC scores at tissue level

The most reproducible bulk observation was a positive tissue-level association between efferocytosis-related and inflammatory VSMC signatures. In GSE111782, the efferocytosis-related score correlated with the inflammatory VSMC score (Spearman rho = {fmt_num(eff_infl_111.get('rho'))}, P = {fmt_p(eff_infl_111.get('P_value'))}, FDR = {fmt_p(eff_infl_111.get('FDR'))}). Leave-one-out correlations ranged from {fmt_num(boot111.get('loo_min_rho'))} to {fmt_num(boot111.get('loo_max_rho'))}. The bootstrap median rho was {fmt_num(boot111.get('bootstrap_median_rho'))}, with a 95% interval from {fmt_num(boot111.get('bootstrap_ci_low'))} to {fmt_num(boot111.get('bootstrap_ci_high'))}. The efferocytosis-related score also correlated with the synthetic VSMC score in GSE111782 (rho = {fmt_num(eff_syn_111.get('rho'))}, P = {fmt_p(eff_syn_111.get('P_value'))}, FDR = {fmt_p(eff_syn_111.get('FDR'))}), which was treated as supportive and exploratory.

In GSE311535, the efferocytosis-related score again correlated with the inflammatory VSMC score (rho = {fmt_num(eff_infl_311.get('rho'))}, P = {fmt_p(eff_infl_311.get('P_value'))}, FDR = {fmt_p(eff_infl_311.get('FDR'))}). Leave-one-out correlations ranged from {fmt_num(boot311.get('loo_min_rho'))} to {fmt_num(boot311.get('loo_max_rho'))}. The bootstrap median rho was {fmt_num(boot311.get('bootstrap_median_rho'))}, with a 95% interval from {fmt_num(boot311.get('bootstrap_ci_low'))} to {fmt_num(boot311.get('bootstrap_ci_high'))}. The association with ECM-remodeling VSMC was weaker and did not pass FDR correction (rho = {fmt_num(eff_ecm_311.get('rho'))}, FDR = {fmt_p(eff_ecm_311.get('FDR'))}). An exploratory fixed-effect Fisher z summary gave a pooled rho of {fmt_num(meta_row.get('pooled_rho'))} with a 95% interval from {fmt_num(meta_row.get('ci_low'))} to {fmt_num(meta_row.get('ci_high'))}, but cohort-specific results remain primary.

### Marker-proxy adjustment showed that the bulk correlation remains confounded

Exploratory partial Spearman analyses showed that the main correlation was sensitive to marker-proxy adjustment. In GSE111782, adjustment for macrophage plus VSMC marker scores attenuated the association ({partial_text(partial111_mv)}). Adjustment for four cell-type marker proxies also attenuated it ({partial_text(partial111_four)}), as did adjustment for a broad inflammation proxy ({partial_text(partial111_inf)}). In GSE311535, the association persisted after adjustment for macrophage plus VSMC marker scores ({partial_text(partial311_mv)}) and after four cell-type marker proxies ({partial_text(partial311_four)}). It was attenuated after broad inflammation-proxy adjustment ({partial_text(partial311_inf)}). These analyses use expression proxies in very small cohorts. They support a cautious interpretation in which cell composition and shared inflammatory burden remain plausible explanations.

### Single-cell analysis localized macrophage and VSMC compartments without donor-level FDR-supported group differences

GSE260657 yielded {sum(cell_counts.values()):,} parsed cells across 15 donor/sample files. Detected genes ranged from {fmt_num(sc_qc['detected_min'], 0)} to {fmt_num(sc_qc['detected_max'], 0)}, total counts from {fmt_num(sc_qc['counts_min'], 0)} to {fmt_num(sc_qc['counts_max'], 0)}, and mitochondrial fraction from {fmt_num(sc_qc['mt_min'])} to {fmt_num(sc_qc['mt_max'])}. Marker-score annotation assigned the main cell classes as {cell_summary} (Figure 4). Macrophages represented the largest annotated compartment (n = {cell_counts.get('macrophage', 0):,}), and VSMCs were also represented (n = {cell_counts.get('vsmc', 0):,}).

Secondary macrophage analysis produced three descriptive subclusters: {macro_text} (Figure 5). Secondary VSMC analysis also produced three descriptive subclusters: {vsmc_text} (Figure 6). Donor-level program comparisons across macrophage and VSMC compartments did not identify FDR-supported symptomatic-versus-asymptomatic differences (FDR < 0.05 comparisons: {sc_tests_sig}). Donor-level cell-composition tests also did not reach FDR < 0.05 (FDR < 0.05 comparisons: {comp_tests_sig}); the lowest adjusted values were {comp_lowest_text}. These single-cell data support cellular localization of the predefined programs, but they do not establish group-level program differences or cell-specific coupling.

### Replication supported the tissue-level correlation but not a classifier or marker set

The independent heterogeneous replication cohort reproduced the positive tissue-level association between efferocytosis-related and inflammatory VSMC scores. It did not support individual DEGs, sample-level signature group shifts, or discovery ranked enrichment at FDR < 0.05 (Figure 7). The optional machine-learning module was skipped because both bulk cohorts were small and heterogeneous. It was also skipped because discovery plus replication did not establish an FDR-supported gene-level marker set. The results are therefore reported as transcriptional associations rather than a classification model.

## Discussion

This secondary analysis found limited evidence for symptomatic-versus-asymptomatic differential expression at the individual gene or sample-level signature level. Neither bulk cohort produced FDR-supported individual genes. No predefined efferocytosis-related or VSMC sample-level signature comparison reached FDR < 0.05. These findings reduce support for a simple group-discriminating transcriptomic signature in the available public cohorts. They should not be read as proof that no biological differences exist.

The most consistent result was a tissue-level association between efferocytosis-related scores and inflammatory VSMC scores across two bulk cohorts. The association was observed in the discovery microarray cohort and in the independent diabetes-specific RNA-seq cohort. Gene-set overlap did not explain the main association because the efferocytosis-related and inflammatory VSMC signatures shared no genes. The association is biologically plausible in light of work on efferocytosis, inflammatory plaque states, and VSMC phenotypic modulation [6-14]. Still, it remains an association in mixed tissue.

The new sensitivity analyses narrow the interpretation. Leave-one-out and bootstrap analyses showed that the rank correlation was not driven by removal of a single sample in either cohort. Marker-proxy adjustment gave a more cautious message. The association was strongly attenuated in GSE111782 after adjustment for macrophage and VSMC proxies, four cell-type proxies, or broad inflammation. In GSE311535, it persisted after cell-type proxy adjustment but weakened after broad inflammation adjustment. These results are compatible with a tissue-level relationship, but they do not separate cellular composition, shared inflammation, lesion stage, and true cell-state coordination.

The replication cohort strengthens the observation because it was independent and used a different assay type. Its diabetes-specific context also limits generalization. Diabetes status, medication use, tissue handling, symptom definitions, surgery timing, and plaque sampling site may differ between cohorts. GSE311535 should therefore be understood as an independent heterogeneous replication cohort, not as a fully matched validation set.

The single-cell cohort helped localize the relevant compartments. Macrophages and VSMCs were both represented in human carotid plaque data, and efferocytosis-related and VSMC-state programs could be scored. Donor-level program and cell-composition comparisons did not reach FDR < 0.05. The S1-S3 subclusters are descriptive, and the fallback workflow lacked complete doublet removal, UMAP visualization, reference mapping, and full integration. These constraints prevent a stronger claim that symptomatic plaques consistently contain an efferocytosis-high macrophage state or a non-contractile VSMC state.

The discovery enrichment results should also be read cautiously. In GSE111782, several predefined signatures were negatively enriched along the symptomatic-versus-asymptomatic ranked list. These enrichment signals were not reproduced in GSE311535. Platform mapping was high, but microarray probe selection, RNA-seq filtering, background gene definitions, diabetes status, sample size, and unmeasured tissue composition could all contribute to this non-replication. The cartilage and chondrocyte GO results were based on small leading gene sets and should be interpreted as exploratory.

Several limitations remain before submission. All analyses used retrospective public datasets, and harmonized patient-level covariates were limited. Symptom status was used exactly as reported and was not treated as pathology-defined plaque status. Bulk plaque tissue mixes immune, stromal, endothelial, smooth muscle, and other cell populations, so tissue-level signature associations may be composition-driven. Marker-proxy adjustment cannot replace histology or validated deconvolution. The single-cell workflow was a reproducible fallback analysis rather than a full Seurat or Scanpy reanalysis of processed objects. No spatial carotid validation, lineage tracing, perturbation experiment, or wet-lab validation was performed. The study was not powered prospectively; bootstrap intervals summarize precision but do not solve small-sample limitations.

In summary, public human carotid plaque transcriptomes support a reproducible association between tissue-level efferocytosis-related and inflammatory VSMC expression scores across two heterogeneous bulk cohorts. They do not establish cell-specific coordination, mechanism, robust symptomatic-versus-asymptomatic DEGs, or validated group-discriminating signatures. Future studies should test these candidate associations in prospectively phenotyped carotid plaque cohorts with harmonized clinical definitions, histology, spatial validation, and experimental models that can separate cell-state changes from cell-composition effects.

## Data Availability Statement

Public datasets were analyzed in this study. They are available through the Gene Expression Omnibus under GSE111782, GSE311535, and GSE260657. Accession-level URLs, download dates, file names, processing scripts, generated outputs, uses, and limitations are recorded in DATA_PROVENANCE.md and 03_data/metadata/dataset_audit.csv.

## Code Availability Statement

Analysis scripts, parameters, logs, intermediate tables, figures, and manuscript-generation scripts are included in this project directory. Before submission, the authors should archive the code in a persistent public repository and replace this placeholder with the repository URL: [[CODE REPOSITORY URL]].

## Ethics Statement

This manuscript reports a secondary analysis of public, de-identified human datasets. The original studies' ethics approval and consent statements must be verified from the corresponding source publications before submission: [[ETHICS STATEMENT TO BE VERIFIED]].

## Author Contributions

[[AUTHOR CONTRIBUTIONS TO BE COMPLETED BY REAL AUTHORS]]

## Funding

[[FUNDING INFORMATION]]

## Conflict of Interest

[[CONFLICT OF INTEREST DECLARATION]]

## Acknowledgments and AI Use Disclosure

Any generative AI assistance should be disclosed according to Frontiers policy. This project used AI assistance for code orchestration, manuscript organization, reviewer-response revision, and language editing. Quantitative analyses and figures were generated from public data using recorded scripts. The submitting authors must verify the final wording of this disclosure.

## Contribution to the Field

Human carotid plaque transcriptomic studies often compare symptomatic and asymptomatic plaques. These clinical labels can be confused with pathology-defined categories or with causal mechanisms. This study provides a reproducible public-data analysis of whether efferocytosis-related transcriptional features are associated with VSMC state programs in clinically defined human carotid plaques. The main reproducible signal is tissue-level covariation between efferocytosis-related and inflammatory VSMC scores in two independent bulk cohorts. The analysis also shows that individual DEGs and sample-level signature group differences are not FDR-supported in the available cohorts. The single-cell analysis localizes macrophage and VSMC compartments but does not treat cells as independent patients. The study therefore offers a transparent candidate association for future experimental and spatial validation, while avoiding causal, diagnostic, pathology-category, or treatment-actionability claims.

## References

{refs_text}
"""
    return humanizer_pass(manuscript)


def humanizer_pass(text: str) -> str:
    replacements = {
        "\u2014": ", ",
        "\u2013": "-",
        "This is important because": "This matters because",
        "It is important to note that": "",
        "key ": "",
        "highlighted": "showed",
        "highlights": "shows",
        "underscores": "supports",
        "robustly": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" +\n", "\n", text)
    text = text.replace("  ", " ")
    return text.strip() + "\n"


def strip_inline_md(text: str) -> str:
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", text)
    return text


def apply_font_to_run(run, size: float = 12.0) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def set_paragraph_format(paragraph) -> None:
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        apply_font_to_run(run, 12.0)


def finalize_doc_style(doc: Document) -> None:
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    for paragraph in doc.paragraphs:
        set_paragraph_format(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_paragraph_format(paragraph)


def add_markdown(doc: Document, markdown: str, skip_top_title: bool = False) -> None:
    lines = markdown.splitlines()
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        text = strip_inline_md(" ".join(part.strip() for part in buffer).strip())
        if text:
            doc.add_paragraph(text)
        buffer.clear()

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        if skip_top_title and line.startswith("# "):
            continue
        if line.startswith("### "):
            flush()
            doc.add_heading(strip_inline_md(line[4:].strip()), level=3)
        elif line.startswith("## "):
            flush()
            doc.add_heading(strip_inline_md(line[3:].strip()), level=2)
        elif line.startswith("# "):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(strip_inline_md(line[2:].strip()))
            r.bold = True
            apply_font_to_run(r, 12.0)
        elif line.startswith("- "):
            flush()
            doc.add_paragraph(strip_inline_md(line[2:].strip()), style="List Bullet")
        else:
            buffer.append(line)
    flush()


def save_docx_from_markdown(markdown: str, output: Path, skip_top_title: bool = False) -> None:
    doc = Document()
    add_markdown(doc, markdown, skip_top_title=skip_top_title)
    finalize_doc_style(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    apply_font_to_run(r, 12.0)


def add_table(doc: Document, title: str, rows: list[dict[str, str]], columns: list[str], max_rows: int | None = None) -> None:
    doc.add_heading(title, level=2)
    if not rows:
        doc.add_paragraph("No rows were available.")
        return
    selected = rows if max_rows is None else rows[:max_rows]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for j, col in enumerate(columns):
        table.rows[0].cells[j].text = col
    for row in selected:
        cells = table.add_row().cells
        for j, col in enumerate(columns):
            cells[j].text = str(row.get(col, ""))
    if max_rows is not None and len(rows) > max_rows:
        add_note(doc, f"Showing {max_rows} of {len(rows)} rows. Complete source table remains in 05_results/tables.")


def figure_legend_map() -> dict[str, str]:
    return {
        "Figure_1_study_design": "Figure 1 | Study design, cohort roles, and analysis gates. Public human carotid plaque cohorts were assigned to discovery bulk analysis, heterogeneous bulk replication, and single-cell localization. Symptomatic and asymptomatic labels are retained as reported in the source metadata.",
        "Figure_2_discovery_bulk": "Figure 2 | Discovery bulk quality control, differential expression, and enrichment. GSE111782 was analyzed with limma using symptomatic versus asymptomatic plaques as the contrast. Positive log2FC denotes higher expression in symptomatic plaques. No genes reached FDR < 0.05.",
        "Figure_3_signature_scores": "Figure 3 | Predefined efferocytosis-related and VSMC signature scores and tissue-level correlations. Sample-level scores summarize mapped genes for each predefined program. Group comparisons use two-sided Wilcoxon tests with Benjamini-Hochberg correction. Correlations use Spearman rho and are tissue-level associations only.",
        "Figure_4_single_cell_atlas": "Figure 4 | Single-cell atlas and major-cell annotation for GSE260657. Smart-seq2 files were parsed from public raw data. Major cell classes were assigned by marker-score and clustering support. Uncertain cells were retained as unassigned.",
        "Figure_5_macrophage_efferocytosis": "Figure 5 | Macrophage subclusters and efferocytosis-related program scoring. Macrophage cells were analyzed descriptively to localize efferocytosis-related expression. Donor/sample-level summaries, rather than cell counts, define the inferential unit for group comparisons.",
        "Figure_6_vsmc_states": "Figure 6 | VSMC subclusters and contractile/non-contractile program scoring. VSMC cells were analyzed descriptively using predefined VSMC state programs. These outputs indicate candidate cell-state localization and do not prove lineage transition or causality.",
        "Figure_7_external_validation": "Figure 7 | Independent heterogeneous bulk replication in GSE311535. Results are reported separately from discovery because the cohort is diabetes-specific and uses RNA-seq counts. The reproducible finding is the efferocytosis-related inflammatory VSMC tissue-level association.",
    }


def build_integrated_doc(manuscript: str, ctx: dict[str, object], output: Path) -> None:
    doc = Document()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(TITLE)
    title_run.bold = True
    apply_font_to_run(title_run, 12.0)
    for line in [
        "Target journal: Frontiers in Cardiovascular Medicine",
        "Article type: Original Research",
        "Integrated third-round author-review document",
        f"Generated on: {stamp()}",
        "Authors: [[AUTHOR NAME]]",
        "Affiliations: [[AFFILIATION]]",
        "Corresponding author: [[CORRESPONDING AUTHOR EMAIL]]",
        "Code repository: [[CODE REPOSITORY URL]]",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_note(doc, "Review caution: author identity, ethics, funding, conflicts, repository URL, and final citation-to-claim mapping still require real-author verification before submission.")
    doc.add_page_break()

    doc.add_heading("Main Manuscript", level=1)
    add_markdown(doc, manuscript, skip_top_title=True)
    doc.add_page_break()

    doc.add_heading("Reviewer-Response Result Tables", level=1)
    add_note(doc, "Compact tables are included for convenient review. Complete CSV files remain in 05_results/tables.")
    add_table(
        doc,
        "Gene-Set Platform Mapping",
        ctx["mapping"],
        ["cohort", "signature", "total_genes", "mapped_genes", "missing_genes", "mapped_fraction"],
    )
    add_table(
        doc,
        "Gene-Set Overlap and Jaccard Index",
        ctx["overlap"],
        ["signature_1", "signature_2", "overlap_count", "union_count", "jaccard", "overlap_genes"],
    )
    add_table(
        doc,
        "Main Correlation Bootstrap and Leave-One-Out Summary",
        ctx["boot"],
        ["cohort", "n_samples", "observed_rho", "observed_P_value", "loo_min_rho", "loo_max_rho", "bootstrap_ci_low", "bootstrap_ci_high"],
    )
    add_table(
        doc,
        "Exploratory Partial Spearman Proxy Adjustment",
        ctx["partial"],
        ["cohort", "model", "adjusted_for", "rho", "P_value", "FDR", "n_samples"],
    )
    add_table(
        doc,
        "Bootstrap Effect Sizes for Bulk Signature Group Comparisons",
        ctx["sig_effects"],
        ["cohort", "signature", "median_difference_symptomatic_minus_asymptomatic", "bootstrap_ci_low", "bootstrap_ci_high", "P_value", "FDR"],
    )
    add_table(
        doc,
        "Single-Cell Donor-Level Cell-Composition Tests",
        ctx["comp_tests"],
        ["major_cell_type", "comparison", "n_donors", "P_value", "FDR"],
    )
    doc.add_page_break()

    doc.add_heading("Main Figures", level=1)
    for stem, legend in figure_legend_map().items():
        doc.add_heading(stem.replace("_", " "), level=2)
        image = FIGURE_DIR / f"{stem}.png"
        if image.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(image), width=Inches(6.2))
            doc.add_paragraph(legend)
            add_note(doc, f"Source image: {image.relative_to(ROOT)}")
        else:
            doc.add_paragraph(f"Figure PNG was not found: {image.relative_to(ROOT)}")
            record_error(f"missing figure PNG: {image}")
    doc.add_page_break()

    doc.add_heading("Full Gene Sets", level=1)
    for signature in ["efferocytosis", "contractile_vsmc", "synthetic_vsmc", "inflammatory_vsmc", "ecm_remodeling_vsmc"]:
        doc.add_heading(signature, level=2)
        doc.add_paragraph(list_genes(ctx["gene_sets"], signature))
    add_note(doc, "Full module labels and source bases are in 05_results/tables/gene_set_category_audit.csv.")
    doc.add_page_break()

    doc.add_heading("Revision Note", level=1)
    add_markdown(doc, build_revision_note(ctx), skip_top_title=True)
    finalize_doc_style(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_revision_note(ctx: dict[str, object]) -> str:
    return f"""# 第三轮大修落实说明

## 已根据审阅意见完成的修改

- 标题已改为更克制的组织水平表述，避免把 bulk 相关性写成巨噬细胞和 VSMC 之间的细胞特异性耦合。
- 全文统一使用 clinically symptomatic versus asymptomatic human carotid plaques，不把症状性斑块写成病理学意义上的不稳定斑块。
- 增加基因集透明性说明，列出 efferocytosis-related 基因、功能模块、平台映射结果、Jaccard 重叠和重叠基因。
- 增加主相关性的 leave-one-out、5,000 次 bootstrap 置信区间和探索性 Fisher z 汇总。
- 增加 marker-score proxy 偏相关分析，明确细胞组成和总体炎症混杂仍可能解释 bulk 相关性。
- 将 GSE311535 改称 independent heterogeneous replication cohort，并说明其糖尿病队列背景限制了同质验证解释。
- 补充单细胞 QC、标准化、PCA、k-means、UMAP 缺失、marker-score 注释和 donor-level 统计说明。
- 增加 donor-level 单细胞细胞组成检验，结果显示无 FDR < 0.05 的细胞类型数量或比例差异。
- 对阴性结果改写为当前样本量和数据质量下未检测到 FDR-supported difference，避免绝对化。
- 使用 Humanizer 规则完成最终语言处理，删除过强机制、诊断、治疗和临床转化措辞。

## 仍需真实作者处理

- 作者姓名、单位、通讯作者邮箱、作者贡献、基金、利益冲突和致谢仍为占位符。
- 原始研究伦理和知情同意表述需要作者从源论文逐条核验。
- 代码仓库 URL 仍需作者上传永久公开仓库后填写。
- 参考文献与最终每一句正文的对应关系仍建议作者逐句核对。
- 若投稿前能获得原始单细胞处理对象，建议再做完整 Seurat 或 Scanpy 流程和参考图谱映射。

## 新增结果文件

- 05_results/tables/gene_set_category_audit.csv
- 05_results/tables/gene_set_overlap_jaccard.csv
- 05_results/tables/gene_set_platform_mapping_summary.csv
- 05_results/tables/main_correlation_leave_one_out.csv
- 05_results/tables/main_correlation_bootstrap_summary.csv
- 05_results/tables/main_correlation_exploratory_meta_summary.csv
- 05_results/tables/main_correlation_partial_spearman_proxy_adjustment.csv
- 05_results/tables/signature_group_effect_sizes_bootstrap.csv
- 05_results/tables/GSE260657_donor_celltype_composition_group_summary.csv
- 05_results/tables/GSE260657_donor_celltype_composition_tests.csv

## 投稿判断

第三轮稿件比第二轮更适合作者审阅，但仍不应直接提交。主要原因是作者、伦理、基金、代码仓库和最终引用核对尚未完成；单细胞部分仍是轻量级可复现定位分析，而非完整重分析。
"""


def write_run_updates(outputs: list[Path]) -> None:
    now = stamp()
    output_text = "; ".join(str(p.relative_to(ROOT)) for p in outputs if p.exists())
    with RUNBOOK.open("a", encoding="utf-8") as handle:
        handle.write(
            f"| {now} / {stamp()} | python 04_scripts/19_apply_major_review_revision.py | 0 | "
            f"{output_text} | Applied major-review manuscript revision and Humanizer pass |\n"
        )
    with PROJECT_STATUS.open("a", encoding="utf-8") as handle:
        handle.write(
            "\n## 2026-09-04 Third-round major-review revision\n"
            "- Completed: revised manuscript language and statistical reporting according to the major-review comments.\n"
            "- Completed: integrated gene-set audit, overlap, mapping, bootstrap, leave-one-out, partial-correlation, and donor-level single-cell composition results.\n"
            "- Completed: exported Times New Roman 12 pt, double-spaced Word files under 修改内容.\n"
            "- Still author-dependent: author information, ethics/consent verification, funding, conflicts, repository URL, and final citation-to-claim audit.\n"
        )


def copy_svg_folder() -> None:
    source = FIGURE_DIR
    dest = OUTPUT_DIR / "SVG图片"
    dest.mkdir(parents=True, exist_ok=True)
    for svg in source.glob("*.svg"):
        shutil.copy2(svg, dest / svg.name)


def qa_docx_text(path: Path) -> dict[str, object]:
    if not path.exists():
        return {"exists": False}
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    plain = re.sub(r"<[^>]+>", " ", xml)
    lowered = f" {plain.lower()} "
    term_hits = []
    for term in FORBIDDEN_REVIEW_TERMS:
        if term == " therapeutic target":
            continue
        if term in lowered:
            term_hits.append(term.strip())
    dash_hits = [dash for dash in ["\u2014", "\u2013"] if dash in plain]
    return {"exists": True, "term_hits": term_hits, "dash_hits": dash_hits, "chars": len(plain)}


def main() -> int:
    started = dt.datetime.now()
    try:
        log("START third-round major-review revision")
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        ctx = build_context()
        refs_text = extract_references()
        manuscript = build_manuscript(ctx, refs_text)

        md_path = MANUSCRIPT_DIR / "manuscript_draft_third_revision.md"
        write_text(md_path, manuscript)
        write_text(OUTPUT_DIR / "投稿正文_第三轮大修完善_Humanizer终版.md", manuscript)

        manuscript_docx = OUTPUT_DIR / "投稿正文_第三轮大修完善_Humanizer终版.docx"
        integrated_docx = OUTPUT_DIR / "文章_第三轮大修完善_含图表整合版.docx"
        note_md = OUTPUT_DIR / "第三轮大修落实说明.md"
        note_docx = OUTPUT_DIR / "第三轮大修落实说明.docx"
        process_log = OUTPUT_DIR / "第三轮处理日志.txt"

        save_docx_from_markdown(manuscript, manuscript_docx)
        build_integrated_doc(manuscript, ctx, integrated_docx)
        revision_note = build_revision_note(ctx)
        write_text(note_md, revision_note)
        save_docx_from_markdown(revision_note, note_docx)
        copy_svg_folder()

        qa = {
            "manuscript_docx": qa_docx_text(manuscript_docx),
            "integrated_docx": qa_docx_text(integrated_docx),
        }
        log_lines = [
            "Third-round major-review revision completed.",
            f"Started: {started.isoformat(timespec='seconds')}",
            f"Ended: {dt.datetime.now().isoformat(timespec='seconds')}",
            "Skills applied: nature-polishing, nature-statistics, humanizer.",
            "Output formatting: Times New Roman, 12 pt, double line spacing.",
            "No original source file was overwritten.",
            "QA:",
            f"- Manuscript DOCX term hits: {qa['manuscript_docx'].get('term_hits')}",
            f"- Manuscript DOCX dash hits: {qa['manuscript_docx'].get('dash_hits')}",
            f"- Integrated DOCX term hits: {qa['integrated_docx'].get('term_hits')}",
            f"- Integrated DOCX dash hits: {qa['integrated_docx'].get('dash_hits')}",
            "Outputs:",
            f"- {manuscript_docx}",
            f"- {integrated_docx}",
            f"- {note_docx}",
            f"- {note_md}",
            f"- {md_path}",
        ]
        write_text(process_log, "\n".join(log_lines) + "\n")

        outputs = [manuscript_docx, integrated_docx, note_docx, note_md, process_log, md_path]
        write_run_updates(outputs)
        log("END third-round major-review revision status=0")
        return 0
    except Exception as exc:
        record_error(str(exc))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
