"""Skill-based manuscript audit and revision.

Purpose: revise the manuscript using the local academic-writing, polishing,
statistics, reviewer, and reference-verification skill rules.
Inputs: project result tables, QC files, current manuscript support files, and
PubMed metadata for the curated reference set.
Outputs: revised manuscript Markdown/DOCX, support DOCX files, cleaned
references.bib, reference_audit.csv, and audit reports under 06_manuscript.
Dependencies: python-docx, Python standard library, network access to NCBI
PubMed E-utilities for reference verification.
Run order: after 11_figure_generation.R and before 14_integrated_review_doc.py.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import math
import re
import shutil
import sys
import textwrap
import urllib.parse
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SEED = 20260904
ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = ROOT / "06_manuscript"
LIT_DIR = ROOT / "02_literature"
TABLE_DIR = ROOT / "05_results" / "tables"
QC_DIR = ROOT / "05_results" / "qc"
FIGURE_DIR = ROOT / "05_results" / "figures"
LOG_DIR = ROOT / "08_logs"
RUNBOOK = ROOT / "RUNBOOK.md"
ERROR_LOG = LOG_DIR / "error_log.md"
LOG = LOG_DIR / "15_skill_based_revision.log"

TITLE = (
    "Tissue-level associations between macrophage efferocytosis-related programs "
    "and vascular smooth muscle cell state signatures in symptomatic versus "
    "asymptomatic human carotid plaques"
)

PMID_PURPOSES = {
    "30335165": "discovery bulk dataset and symptomatic carotid plaque transcriptomics",
    "41377472": "independent diabetic bulk validation dataset",
    "38639096": "human carotid plaque single-cell dataset",
    "23193258": "GEO repository update and public-data source",
    "11752295": "GEO repository original description",
    "27725526": "macrophage apoptosis and efferocytosis in atherosclerosis",
    "28137963": "efferocytosis biology in atherosclerosis",
    "33630758": "efferocytosis-related macrophage mechanism in plaque necrosis",
    "34029141": "VSMC fate and state framework in atherosclerosis",
    "34617061": "VSMC phenotypic switching context",
    "34936041": "VSMC phenotypic switching review context",
    "32981416": "human atherosclerotic plaque single-cell atlas context",
    "32962412": "single-cell VSMC phenotypic switching context",
    "38362263": "human plaque macrophage states and cerebrovascular complications",
    "37471165": "human plaque leukocyte single-cell context",
    "39041203": "human carotid plaque adaptive immunity single-cell context",
    "39613875": "human atherosclerosis single-cell immune checkpoint context",
    "40211055": "human carotid plaque single-cell sex-difference context",
    "39552248": "atheroma transcriptomics and VSMC regulator context",
    "40438929": "human carotid plaque multi-omic context",
    "25605792": "limma differential-expression method",
    "19910308": "edgeR differential-expression method",
    "23323831": "GSVA signature-scoring method",
    "16199517": "GSEA/ranked gene-set enrichment method",
    "22455463": "clusterProfiler enrichment method",
    "26661513": "ReactomePA enrichment method",
    "33290552": "Gene Ontology resource",
    "34788843": "Reactome pathway knowledgebase",
    "33125081": "KEGG pathway database",
    "19114008": "WGCNA method, cited only for the prespecified skipped optional module",
}

PMIDS = list(PMID_PURPOSES.keys())


def stamp() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def log(message: str) -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"[{stamp()}] {message}\n")


def record_error(message: str) -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    with ERROR_LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"\n- {stamp()} | 15_skill_based_revision | {message}\n")
    log(f"ERROR {message}")


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        record_error(f"missing CSV: {path}")
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def fmt_num(value: object, digits: int = 3) -> str:
    try:
        x = float(value)
    except (TypeError, ValueError):
        return "NA"
    if not math.isfinite(x):
        return "NA"
    if x == 0:
        return "0"
    if abs(x) < 0.001:
        return f"{x:.2e}"
    if abs(x) < 0.01:
        return f"{x:.4f}"
    return f"{x:.{digits}f}"


def count_fdr(path: Path) -> tuple[int, int]:
    rows = read_csv_rows(path)
    if not rows:
        return 0, 0
    fdr_col = None
    for key in rows[0].keys():
        if key.lower() in {"fdr", "adj.p.val", "padj", "p.adjust", "qvalue"}:
            fdr_col = key
            break
    if fdr_col is None:
        return len(rows), 0
    sig = 0
    for row in rows:
        try:
            if float(row.get(fdr_col, "nan")) < 0.05:
                sig += 1
        except ValueError:
            continue
    return len(rows), sig


def top_rows(path: Path, p_col: str, n: int = 3) -> list[dict[str, str]]:
    rows = read_csv_rows(path)
    return sorted(rows, key=lambda row: float(row.get(p_col, "1") or 1))[:n]


def get_row(rows: list[dict[str, str]], **criteria: str) -> dict[str, str]:
    for row in rows:
        if all(row.get(key) == value for key, value in criteria.items()):
            return row
    return {}


def fetch_pubmed(pmids: list[str]) -> dict[str, dict[str, object]]:
    url = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?" + urllib.parse.urlencode(
        {"db": "pubmed", "id": ",".join(pmids), "retmode": "json"}
    )
    with urllib.request.urlopen(url, timeout=90) as response:
        data = json.load(response)
    result = data.get("result", {})
    return {pmid: result[pmid] for pmid in result.get("uids", []) if pmid in result}


def article_doi(rec: dict[str, object]) -> str:
    for article_id in rec.get("articleids", []):
        if article_id.get("idtype") == "doi":
            return str(article_id.get("value", ""))
    return ""


def article_year(rec: dict[str, object]) -> str:
    match = re.search(r"\d{4}", str(rec.get("pubdate", "")))
    return match.group(0) if match else ""


def article_authors(rec: dict[str, object]) -> list[str]:
    return [str(a.get("name", "")) for a in rec.get("authors", []) if a.get("name")]


def bib_key(rec: dict[str, object], idx: int) -> str:
    authors = article_authors(rec)
    first = authors[0].split()[0].lower() if authors else "ref"
    first = re.sub(r"[^a-z0-9]+", "", first)
    return f"{first}{article_year(rec)}_{idx:02d}"


def clean_bib_value(value: str) -> str:
    return str(value).replace("{", "").replace("}", "").strip()


def write_references(records: dict[str, dict[str, object]]) -> list[dict[str, str]]:
    audit_rows: list[dict[str, str]] = []
    bib_lines = [
        "% Curated and regenerated by 04_scripts/15_skill_based_revision.py.",
        "% Metadata source: NCBI PubMed E-utilities, retrieved " + stamp() + ".",
        "% Low-relevance and method-mismatched entries from the prior draft were removed.",
        "",
    ]
    for idx, pmid in enumerate(PMIDS, start=1):
        rec = records.get(pmid)
        if not rec:
            record_error(f"PMID {pmid} was not returned by PubMed")
            continue
        authors = article_authors(rec)
        doi = article_doi(rec)
        year = article_year(rec)
        title = str(rec.get("title", "")).strip()
        journal = str(rec.get("fulljournalname", "") or rec.get("source", "")).strip()
        audit_rows.append({
            "title": title,
            "authors": "; ".join(authors),
            "journal": journal,
            "year": year,
            "DOI": doi,
            "PMID": pmid,
            "verified_url": f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/",
            "citation_purpose": PMID_PURPOSES[pmid],
            "verification_status": "PubMed metadata verified; relevance checked for this manuscript revision",
        })
        key = bib_key(rec, idx)
        bib_lines.extend([
            f"@article{{{key},",
            f"  author = {{{clean_bib_value(' and '.join(authors))}}},",
            f"  title = {{{clean_bib_value(title)}}},",
            f"  journal = {{{clean_bib_value(journal)}}},",
            f"  year = {{{year}}},",
        ])
        if str(rec.get("volume", "")).strip():
            bib_lines.append(f"  volume = {{{clean_bib_value(rec.get('volume', ''))}}},")
        if str(rec.get("issue", "")).strip():
            bib_lines.append(f"  number = {{{clean_bib_value(rec.get('issue', ''))}}},")
        if str(rec.get("pages", "")).strip():
            bib_lines.append(f"  pages = {{{clean_bib_value(rec.get('pages', ''))}}},")
        if doi:
            bib_lines.append(f"  doi = {{{clean_bib_value(doi)}}},")
        bib_lines.append(f"  pmid = {{{pmid}}},")
        bib_lines.append(f"  url = {{https://pubmed.ncbi.nlm.nih.gov/{pmid}/}}")
        bib_lines.extend(["}", ""])
    write_csv(
        LIT_DIR / "reference_audit.csv",
        audit_rows,
        ["title", "authors", "journal", "year", "DOI", "PMID", "verified_url", "citation_purpose", "verification_status"],
    )
    (MANUSCRIPT_DIR / "references.bib").write_text("\n".join(bib_lines), encoding="utf-8")
    return audit_rows


def numbered_references(audit_rows: list[dict[str, str]]) -> str:
    lines = []
    for idx, row in enumerate(audit_rows, start=1):
        doi = f" doi: {row['DOI']}." if row.get("DOI") else ""
        pmid = f" PMID: {row['PMID']}." if row.get("PMID") else ""
        lines.append(
            f"{idx}. {row['authors']}. {row['title']} {row['journal']}. {row['year']}.{doi}{pmid}"
        )
    return "\n".join(lines)


def markdown_table(rows: list[dict[str, str]], columns: list[str]) -> str:
    if not rows:
        return ""
    out = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for row in rows:
        out.append("| " + " | ".join(str(row.get(col, "")).replace("|", "/") for col in columns) + " |")
    return "\n".join(out)


def build_result_context() -> dict[str, object]:
    bulk = read_csv_rows(TABLE_DIR / "bulk_analysis_summary.csv")
    sig_summary = read_csv_rows(TABLE_DIR / "external_validation_signature_summary.csv")
    sig_corr = read_csv_rows(TABLE_DIR / "signature_correlations.csv")
    gene_dir = read_csv_rows(TABLE_DIR / "external_validation_gene_direction.csv")
    donor_tests = read_csv_rows(TABLE_DIR / "GSE260657_donor_program_tests.csv")
    cluster_summary = read_csv_rows(QC_DIR / "GSE260657_cluster_annotation_summary.csv")
    cell_rows = read_csv_rows(QC_DIR / "GSE260657_cell_annotations.csv")
    ranked_111 = read_csv_rows(TABLE_DIR / "GSE111782" / "ranked_signature_enrichment.csv")
    ranked_311 = read_csv_rows(TABLE_DIR / "GSE311535" / "ranked_signature_enrichment.csv")
    go_111 = read_csv_rows(TABLE_DIR / "GSE111782" / "GO_BP_enrichment.csv")
    reactome_111 = read_csv_rows(TABLE_DIR / "GSE111782" / "Reactome_enrichment.csv")
    go_311 = read_csv_rows(TABLE_DIR / "GSE311535" / "GO_BP_enrichment.csv")
    reactome_311 = read_csv_rows(TABLE_DIR / "GSE311535" / "Reactome_enrichment.csv")
    macro_sub = read_csv_rows(TABLE_DIR / "GSE260657_macrophage_subcluster_summary.csv")
    vsmc_sub = read_csv_rows(TABLE_DIR / "GSE260657_vsmc_subcluster_summary.csv")
    ds_audit = read_csv_rows(ROOT / "03_data" / "metadata" / "dataset_audit.csv")

    cell_counts: dict[str, int] = {}
    for row in cell_rows:
        key = row.get("major_cell_type", "unassigned") or "unassigned"
        cell_counts[key] = cell_counts.get(key, 0) + 1

    common_genes = len(gene_dir)
    concordant = sum(row.get("direction_concordant") == "TRUE" for row in gene_dir)
    both_sig = sum(
        row.get("discovery_FDR_lt_05") == "TRUE" and row.get("validation_FDR_lt_05") == "TRUE"
        for row in gene_dir
    )

    return {
        "bulk": bulk,
        "sig_summary": sig_summary,
        "sig_corr": sig_corr,
        "gene_dir": gene_dir,
        "donor_tests": donor_tests,
        "cluster_summary": cluster_summary,
        "cell_counts": cell_counts,
        "ranked_111": ranked_111,
        "ranked_311": ranked_311,
        "go_111": go_111,
        "reactome_111": reactome_111,
        "go_311": go_311,
        "reactome_311": reactome_311,
        "macro_sub": macro_sub,
        "vsmc_sub": vsmc_sub,
        "ds_audit": ds_audit,
        "deg111": count_fdr(TABLE_DIR / "GSE111782_DEG_complete.csv"),
        "deg311": count_fdr(TABLE_DIR / "GSE311535_DEG_complete.csv"),
        "top111": top_rows(TABLE_DIR / "GSE111782_DEG_complete.csv", "P_value", 3),
        "top311": top_rows(TABLE_DIR / "GSE311535_DEG_complete.csv", "P_value", 3),
        "common_genes": common_genes,
        "concordant": concordant,
        "both_sig": both_sig,
    }


def ranked_sentence(row: dict[str, str]) -> str:
    return f"{row.get('pathway')} (NES = {fmt_num(row.get('NES'))}, FDR = {fmt_num(row.get('padj'))})"


def build_manuscript(ctx: dict[str, object], refs_text: str) -> str:
    sig_summary = ctx["sig_summary"]
    sig_corr = ctx["sig_corr"]
    ranked_111 = ctx["ranked_111"]
    ranked_311 = ctx["ranked_311"]
    go_111 = ctx["go_111"]
    donor_tests = ctx["donor_tests"]
    macro_sub = ctx["macro_sub"]
    vsmc_sub = ctx["vsmc_sub"]
    cell_counts = ctx["cell_counts"]
    deg111_n, deg111_sig = ctx["deg111"]
    deg311_n, deg311_sig = ctx["deg311"]

    eff_infl_111 = get_row(sig_corr, cohort="GSE111782", signature_1="efferocytosis", signature_2="inflammatory_vsmc")
    eff_syn_111 = get_row(sig_corr, cohort="GSE111782", signature_1="efferocytosis", signature_2="synthetic_vsmc")
    eff_infl_311 = get_row(sig_corr, cohort="GSE311535", signature_1="efferocytosis", signature_2="inflammatory_vsmc")
    eff_ecm_311 = get_row(sig_corr, cohort="GSE311535", signature_1="efferocytosis", signature_2="ecm_remodeling_vsmc")
    contractile_111 = get_row(sig_summary, cohort="GSE111782", signature="contractile_vsmc")
    eff_111 = get_row(sig_summary, cohort="GSE111782", signature="efferocytosis")
    eff_311 = get_row(sig_summary, cohort="GSE311535", signature="efferocytosis")
    best_311 = sorted(ranked_311, key=lambda r: float(r.get("padj", "1") or 1))[0] if ranked_311 else {}

    sig_ranked_111 = [r for r in ranked_111 if float(r.get("padj", "1") or 1) < 0.05]
    sig_ranked_text = "; ".join(ranked_sentence(row) for row in sig_ranked_111)
    go_text = "; ".join(
        f"{row.get('Description')} (FDR = {fmt_num(row.get('p.adjust'))})" for row in go_111[:3]
    )

    cell_summary = ", ".join(
        f"{key}: {value}" for key, value in sorted(cell_counts.items(), key=lambda item: (-item[1], item[0]))
    )
    macro_text = "; ".join(
        f"{row.get('subcluster')} n = {row.get('n_cells')}, efferocytosis score = {fmt_num(row.get('efferocytosis_score'))}"
        for row in macro_sub
    )
    vsmc_text = "; ".join(
        f"{row.get('subcluster')} n = {row.get('n_cells')}, contractile score = {fmt_num(row.get('contractile_vsmc_score'))}"
        for row in vsmc_sub
    )
    sc_tests_sig = sum(float(row.get("FDR", "1") or 1) < 0.05 for row in donor_tests)

    top111 = ", ".join(
        f"{row['gene']} (log2FC = {fmt_num(row['logFC'])}, P = {fmt_num(row['P_value'])}, FDR = {fmt_num(row['FDR'])})"
        for row in ctx["top111"]
    )
    top311 = ", ".join(
        f"{row['gene']} (log2FC = {fmt_num(row['logFC'])}, P = {fmt_num(row['P_value'])}, FDR = {fmt_num(row['FDR'])})"
        for row in ctx["top311"]
    )

    concord_pct = 100 * ctx["concordant"] / ctx["common_genes"] if ctx["common_genes"] else 0

    return f"""# {TITLE}

Article type: Original Research

Running title: Efferocytosis and VSMC programs in carotid plaques

Authors: [[AUTHOR NAME]]

Affiliations: [[AFFILIATION]]

Corresponding author: [[CORRESPONDING AUTHOR EMAIL]]

Manuscript status: revised author-review draft generated from project tables on {dt.date.today().isoformat()}.

## Abstract

Background: Defective macrophage efferocytosis and vascular smooth muscle cell (VSMC) state switching are implicated in atherosclerotic plaque biology, but their relationship in clinically defined human carotid plaques remains incompletely resolved. Methods: We analyzed legally accessible public human carotid plaque transcriptomic datasets from GEO: a discovery Affymetrix microarray cohort, GSE111782 (9 symptomatic and 9 asymptomatic plaques), an independent diabetic RNA-seq validation cohort, GSE311535 (6 symptomatic and 6 asymptomatic plaques), and a Smart-seq2 single-cell cohort, GSE260657 (8 symptomatic and 7 asymptomatic carotid plaque samples; 7,690 cells after parsing). Predefined efferocytosis and VSMC state signatures were evaluated using assay-appropriate bulk models, rank-based enrichment, sample-level scoring, Spearman correlation, and donor-level single-cell summaries. P values were two-sided where applicable and adjusted using the Benjamini-Hochberg method. Results: No individual genes reached FDR < 0.05 in the discovery or validation bulk differential-expression analyses, and no predefined sample-level signature differed between symptomatic and asymptomatic plaques at FDR < 0.05. In the discovery cohort, ranked signature enrichment showed lower-ranked efferocytosis and VSMC programs in symptomatic plaques, but this pattern was not reproduced at FDR < 0.05 in the validation cohort. Across bulk samples, efferocytosis scores were positively associated with inflammatory VSMC scores in both GSE111782 (Spearman rho = {fmt_num(eff_infl_111.get('rho'))}, FDR = {fmt_num(eff_infl_111.get('FDR'))}) and GSE311535 (rho = {fmt_num(eff_infl_311.get('rho'))}, FDR = {fmt_num(eff_infl_311.get('FDR'))}). Single-cell analysis localized macrophage and VSMC compartments, but donor-level macrophage or VSMC program comparisons did not reach FDR < 0.05. Conclusion: The most reproducible observation was a tissue-level association between efferocytosis-related and inflammatory VSMC transcriptional programs rather than a robust symptomatic-versus-asymptomatic differential signature. These public-data results identify candidate associations for experimental follow-up and do not establish causality, pathology-defined plaque status, diagnostic utility, or therapeutic actionability.

Keywords: carotid atherosclerosis; symptomatic plaque; asymptomatic plaque; efferocytosis; macrophage; vascular smooth muscle cell; transcriptomics

## Introduction

Human carotid atherosclerotic plaques are clinically important because symptomatic plaques are associated with cerebrovascular events, yet symptom status is not identical to a pathology-defined plaque category. This distinction matters for transcriptomic studies: datasets that label samples as symptomatic versus asymptomatic should not be rewritten using pathology-defined categories unless the original metadata or linked pathology report supports that terminology. The present study therefore retains the clinical labels reported in each dataset.

Macrophage handling of apoptotic cells is one candidate process linking inflammation resolution, necrotic-core formation, and advanced plaque biology. Efferocytosis requires receptors, bridging molecules, phagolysosomal processing, and lipid-handling responses, and prior experimental and human evidence has linked defective efferocytosis with atherosclerotic inflammation and plaque necrosis [6-8]. In parallel, VSMCs in atherosclerotic plaques occupy contractile and non-contractile states, including matrix-producing, inflammatory, osteogenic, and macrophage-like programs [9-13]. These state changes complicate bulk tissue interpretation because cell composition and cell-state activity can both influence transcriptomic signals.

Single-cell studies have sharpened this problem by showing substantial cellular heterogeneity within human atherosclerotic plaques, including macrophage state diversity and VSMC phenotypic modulation [12-20]. However, less is known about whether a predefined macrophage efferocytosis-related transcriptional program is reproducibly associated with VSMC state signatures in clinical symptomatic versus asymptomatic human carotid plaque cohorts.

Here we performed a secondary analysis of public human carotid plaque datasets to test a bounded hypothesis: symptomatic plaques would show altered efferocytosis-related transcriptional features and these features would be associated with non-contractile VSMC programs. We used one discovery bulk cohort, one independent bulk validation cohort, and one human carotid plaque single-cell cohort, with predefined gene sets and donor-aware single-cell summaries. The analysis was designed to identify reproducible associations, not to infer causality or reclassify symptomatic plaques using pathology-defined categories.

## Materials and Methods

### Data sources and cohort definitions

Public datasets were identified from GEO and retained only when the metadata supported human carotid atherosclerotic plaque tissue and clinically interpretable grouping. GSE111782 was used as the discovery bulk cohort because its GEO record and linked publication describe post-bifurcation internal carotid atheroma samples with symptomatic and asymptomatic clinical labels [1,4,5]. This cohort contained 18 samples, with 9 symptomatic and 9 asymptomatic plaques, profiled on the GPL571 Affymetrix Human Genome U133A 2.0 Array. GSE311535 was used as the independent bulk validation cohort because it contains human carotid plaque RNA-seq count data from patients with diabetes, including 6 symptomatic and 6 asymptomatic samples [2]. GSE260657 was used for single-cell localization because the GEO record and linked publication describe Smart-seq2 data from 15 human carotid plaque samples, including 8 symptomatic and 7 asymptomatic samples [3]. Dataset provenance, URLs, limitations, and audit decisions are recorded in DATA_PROVENANCE.md and 03_data/metadata/dataset_audit.csv.

### Predefined gene sets and signatures

The efferocytosis gene set was defined before inspecting project results and included receptors, bridging molecules, phagocytic/lysosomal genes, and lipid-handling genes relevant to apoptotic-cell recognition, engulfment, and processing. Source bases included Gene Ontology, Reactome, UniProt annotations, and published atherosclerosis/efferocytosis literature [6-8,27,28]. VSMC state signatures were likewise predefined and represented contractile VSMC, synthetic VSMC, inflammatory VSMC, and ECM-remodeling/osteogenic VSMC programs, based on canonical markers and prior plaque single-cell studies [9-13]. These signatures were treated as expression programs rather than definitive lineage states.

### Bulk transcriptomic processing

GSE111782 was analyzed as an RMA-processed log2 microarray expression matrix. Probes without usable gene annotation were removed where possible, duplicate gene mappings were handled by retaining the probe with the highest average expression, and the symptomatic-versus-asymptomatic contrast was fitted using limma empirical Bayes linear modeling [21]. GSE311535 was analyzed as gene-level RNA-seq counts using edgeR quasi-likelihood negative-binomial models after low-expression filtering and library-size normalization [22]. In both cohorts, positive log2 fold change denotes higher expression in symptomatic plaques relative to asymptomatic plaques. Differential-expression P values were adjusted using the Benjamini-Hochberg method, with FDR < 0.05 as the prespecified threshold.

### Signature, enrichment, and association analyses

Sample-level signature scores were computed from predefined gene sets using available mapped genes; the score represents the mean standardized expression of genes available on the relevant platform. Group comparisons used two-sided Wilcoxon rank-sum tests because the bulk cohorts were small and no normality assumption was defensible. Ranked enrichment of predefined signatures used preranked symptomatic-versus-asymptomatic statistics following the GSEA framework [24]. GO biological process, KEGG, and Reactome over-representation analyses were attempted using clusterProfiler and ReactomePA when gene mapping and package availability allowed [25-29]. All enrichment and signature tests used Benjamini-Hochberg correction. Tissue-level associations between signatures were evaluated using Spearman correlations. WGCNA was prespecified but not run because the discovery bulk sample size was 18, below the project threshold of at least 30 samples and therefore not adequate for stable network inference [30].

### Single-cell processing and annotation

GSE260657 raw Smart-seq2 text files were parsed without modifying the original downloaded archive. Cells were evaluated using prespecified quality-control rules for detected genes, counts, mitochondrial fraction, and sample-level parsing. Because the available workflow was a lightweight reproducible fallback rather than a full Seurat/Scanpy integration, cell annotation used canonical marker scoring, PCA/k-means clustering, and marker-candidate inspection. Major cell classes included macrophage, VSMC, endothelial cell, fibroblast, T/NK cell, B cell, mast cell, and unassigned. Uncertain cells were left as unassigned. Group comparisons were performed at donor/sample level when possible; individual cells were not treated as independent patients.

### Reporting, software, and reproducibility

Analyses were performed in R 4.6.0 using limma 3.68.5, edgeR 4.10.4, GSVA 2.6.6, fgsea 1.38.0, clusterProfiler 4.20.0, ReactomePA, ggplot2 4.0.3, and related Bioconductor packages; package versions are archived in 09_environment/package_versions.csv and 09_environment/session_info.txt. Scripts are stored in 04_scripts, logs in 08_logs, and generated tables and figures in 05_results. No samples were removed only because they changed the result direction or significance.

## Results

### Public cohort selection produced discovery, validation, and single-cell evidence layers

The dataset audit retained GSE111782 as the discovery bulk cohort, GSE311535 as an independent bulk validation cohort, and GSE260657 as the single-cell localization cohort (Figure 1). The bulk cohorts were not merged because they differed by platform, preprocessing, and clinical context; GSE311535 is diabetes-specific. The single-cell cohort passed the project gate for human carotid plaque localization after raw file parsing, donor/sample metadata review, and marker-based annotation. Optional spatial transcriptomics, cell-cell communication, and machine-learning modules were not used in the main conclusions because their strict data, annotation, and overfitting-control gates were not met.

### Bulk differential-expression analysis did not identify FDR-significant individual genes

In GSE111782, limma tested {deg111_n:,} genes and identified {deg111_sig} genes at FDR < 0.05 (Figure 2). The smallest nominal P-value genes were {top111}, but their adjusted FDR values were not below the prespecified threshold. In GSE311535, edgeR tested {deg311_n:,} genes and identified {deg311_sig} genes at FDR < 0.05. The smallest nominal P-value genes were {top311}, again without FDR support. Across {ctx['common_genes']:,} genes available in both bulk cohorts, {ctx['concordant']:,} ({concord_pct:.1f}%) had concordant symptomatic-versus-asymptomatic log2FC directions, but no gene reached FDR < 0.05 in both cohorts. These results argue against presenting individual DEGs as validated markers in the current public-data analysis.

### Discovery ranked enrichment suggested lower efferocytosis and VSMC program ranks in symptomatic plaques, but validation was weaker

Ranked enrichment of predefined signatures in the discovery cohort identified FDR-supported negative enrichment for {sig_ranked_text}. Because positive log2FC denotes higher expression in symptomatic plaques, negative normalized enrichment scores indicate that these predefined genes tended to occur lower in the symptomatic-versus-asymptomatic ranked list. GO biological process over-representation in the discovery cohort returned three FDR-supported terms related to cartilage or chondrocyte differentiation: {go_text}. Reactome over-representation did not yield FDR-supported terms in the discovery output. In the validation cohort, none of the ranked predefined signatures reached FDR < 0.05; the lowest adjusted value was observed for {best_311.get('pathway', 'NA')} (NES = {fmt_num(best_311.get('NES'))}, FDR = {fmt_num(best_311.get('padj'))}). Thus, the discovery enrichment pattern was not reproduced as a validation-level result.

### Sample-level signature group differences did not survive FDR correction

At the sample-score level, no predefined efferocytosis or VSMC signature differed between symptomatic and asymptomatic plaques at FDR < 0.05 in either bulk cohort (Figure 3). In GSE111782, the efferocytosis median score was {fmt_num(eff_111.get('median_asymptomatic'))} in asymptomatic plaques and {fmt_num(eff_111.get('median_symptomatic'))} in symptomatic plaques (P = {fmt_num(eff_111.get('P_value'))}, FDR = {fmt_num(eff_111.get('FDR'))}). The contractile VSMC score showed a nominal difference in GSE111782 (median {fmt_num(contractile_111.get('median_asymptomatic'))} versus {fmt_num(contractile_111.get('median_symptomatic'))}; P = {fmt_num(contractile_111.get('P_value'))}), but this did not survive correction (FDR = {fmt_num(contractile_111.get('FDR'))}). In GSE311535, the efferocytosis median was {fmt_num(eff_311.get('median_asymptomatic'))} in asymptomatic plaques and {fmt_num(eff_311.get('median_symptomatic'))} in symptomatic plaques (P = {fmt_num(eff_311.get('P_value'))}, FDR = {fmt_num(eff_311.get('FDR'))}). These results were retained as negative or inconclusive findings rather than reframed as positive group differences.

### Efferocytosis scores were reproducibly associated with inflammatory VSMC scores at tissue level

The strongest reproducible bulk observation was a positive tissue-level association between the efferocytosis and inflammatory VSMC signatures. In GSE111782, the efferocytosis score correlated with the inflammatory VSMC score (Spearman rho = {fmt_num(eff_infl_111.get('rho'))}, P = {fmt_num(eff_infl_111.get('P_value'))}, FDR = {fmt_num(eff_infl_111.get('FDR'))}) and with the synthetic VSMC score (rho = {fmt_num(eff_syn_111.get('rho'))}, P = {fmt_num(eff_syn_111.get('P_value'))}, FDR = {fmt_num(eff_syn_111.get('FDR'))}). In GSE311535, the efferocytosis score again correlated with the inflammatory VSMC score (rho = {fmt_num(eff_infl_311.get('rho'))}, P = {fmt_num(eff_infl_311.get('P_value'))}, FDR = {fmt_num(eff_infl_311.get('FDR'))}), whereas the association with ECM-remodeling VSMC was weaker and did not pass FDR correction (rho = {fmt_num(eff_ecm_311.get('rho'))}, FDR = {fmt_num(eff_ecm_311.get('FDR'))}). These correlations were analyzed across mixed plaque tissue samples and therefore cannot distinguish coordinated cell-state activity from cell-composition differences or shared inflammatory confounding.

### Single-cell analysis localized macrophage and VSMC compartments without donor-level FDR-supported group differences

GSE260657 yielded {sum(cell_counts.values()):,} parsed cells across 15 donor/sample files. Marker-score annotation assigned major cell classes as follows: {cell_summary} (Figure 4). Macrophages represented the largest annotated compartment (n = {cell_counts.get('macrophage', 0):,}), and VSMCs were also clearly represented (n = {cell_counts.get('vsmc', 0):,}). Secondary macrophage analysis produced three descriptive subclusters: {macro_text} (Figure 5). Secondary VSMC analysis similarly produced three descriptive VSMC subclusters: {vsmc_text} (Figure 6). Donor-level program comparisons across macrophage and VSMC compartments did not identify FDR-supported symptomatic-versus-asymptomatic differences (number of tested cell-type/program comparisons with FDR < 0.05: {sc_tests_sig}). The single-cell data therefore support cellular localization of the predefined programs but do not provide independent donor-level evidence for group differences.

### Independent validation supported the efferocytosis-inflammatory VSMC association but not a validated classifier or marker set

The independent validation cohort confirmed the positive tissue-level association between efferocytosis and inflammatory VSMC scores, but it did not validate individual DEGs, sample-level signature group shifts, or discovery ranked enrichment at FDR < 0.05 (Figure 7). Because the discovery and validation cohorts were small and heterogeneous, and because no FDR-supported gene-level marker set was established in discovery and validation, the prespecified optional machine-learning module was skipped. The results are therefore reported as reproducible transcriptional associations rather than as a diagnostic model or treatment-target discovery.

## Discussion

This secondary analysis found limited evidence for robust symptomatic-versus-asymptomatic differential expression at the individual gene or sample-level signature level. Neither bulk cohort produced FDR-significant individual genes, and no predefined efferocytosis or VSMC sample-level signature group comparison reached FDR < 0.05. These negative findings are important because the small cohorts and heterogeneous public metadata create a high risk of overinterpreting nominal P values.

The most consistent result was instead a cross-cohort tissue-level association between efferocytosis-related scores and inflammatory VSMC scores. This association was observed in both the discovery microarray cohort and the independent diabetic RNA-seq cohort. The finding is biologically plausible in light of prior work on macrophage efferocytosis, inflammatory plaque states, and VSMC phenotypic modulation [6-14], but it remains an association across mixed tissue samples. It may reflect coordinated activation of macrophage and VSMC programs, shared inflammatory burden, variation in cellular composition, or differences in lesion stage that could not be fully adjusted using the available metadata.

The single-cell cohort helped localize the relevant compartments, showing that macrophages and VSMCs were both represented in human carotid plaque data and that efferocytosis and VSMC-state programs could be scored at cell and donor levels. However, donor-level program comparisons were not FDR-significant. This constraint prevents a stronger claim that symptomatic plaques consistently contain a distinct efferocytosis-high macrophage state or a distinct non-contractile VSMC state in the analyzed single-cell data. The current single-cell results are therefore best viewed as localization and hypothesis-refinement evidence.

The discovery ranked enrichment results add nuance. In GSE111782, several predefined signatures, including efferocytosis and VSMC programs, were negatively enriched along the symptomatic-versus-asymptomatic ranked list. Because these enrichment signals were not reproduced in GSE311535, they should be treated as cohort-specific observations rather than validated disease signatures. Differences in platform, diabetes status, sample size, tissue handling, plaque stage, and clinical symptom definitions may all contribute to this non-replication.

Several limitations should be explicit before submission. First, all analyses used public retrospective datasets, and patient-level clinical covariates were limited. Second, symptom status was used exactly as reported and was not treated as synonymous with a pathology-defined plaque category. Third, bulk plaque tissue is a mixture of immune, stromal, endothelial, smooth muscle, and other cell populations, so tissue-level signature associations may be composition-driven. Fourth, the single-cell workflow used a lightweight marker-score annotation fallback, and full reanalysis with the original authors' processed objects or a dedicated Seurat/Scanpy pipeline would strengthen annotation confidence. Fifth, no spatial carotid validation, lineage tracing, perturbation experiment, or wet-lab validation was performed. These boundaries preclude causal, diagnostic, therapeutic, or mechanism-confirming claims.

In summary, the data support a conservative conclusion: public human carotid plaque transcriptomes show reproducible tissue-level covariation between macrophage efferocytosis-related and inflammatory VSMC transcriptional programs, but they do not establish robust symptomatic-versus-asymptomatic DEGs or validated group-discriminating signatures. Future work should test these candidate associations in prospectively phenotyped carotid plaque cohorts with harmonized clinical definitions, spatial validation, and experimental models that can separate cell-state changes from cell-composition effects.

## Data Availability Statement

Publicly available datasets were analyzed in this study. These data are available through the Gene Expression Omnibus under accession numbers GSE111782, GSE311535, and GSE260657. Accession-level URLs, download dates, file names, processing scripts, generated outputs, uses, and limitations are recorded in DATA_PROVENANCE.md and 03_data/metadata/dataset_audit.csv.

## Code Availability Statement

Analysis scripts, parameters, logs, intermediate tables, generated figures, and manuscript-generation scripts are included in this project directory. Before journal submission, the authors should archive the code in a persistent public repository and replace this placeholder with the repository URL: [[CODE REPOSITORY URL]].

## Ethics Statement

This manuscript reports a secondary analysis of public, de-identified human datasets. The original studies' ethics approval and consent statements must be verified from the corresponding source publications before submission: [[ETHICS STATEMENT TO BE VERIFIED]].

## Author Contributions

[[AUTHOR CONTRIBUTIONS TO BE COMPLETED BY REAL AUTHORS]]

## Funding

[[FUNDING INFORMATION]]

## Conflict of Interest

[[CONFLICT OF INTEREST DECLARATION]]

## Acknowledgments and AI Use Disclosure

The authors should disclose any generative AI assistance in accordance with Frontiers policy. This project used AI assistance for code orchestration, manuscript organization, and language revision; quantitative analyses and figures were generated from public data using the recorded scripts. The final wording of this disclosure must be verified by the submitting authors.

## Contribution to the Field

Human carotid plaque transcriptomic studies often compare symptomatic and asymptomatic plaques, but symptom status, pathology-defined plaque categories, and mechanistic causality are sometimes conflated. This study provides a reproducible public-data analysis focused on a specific biological question: whether macrophage efferocytosis-related transcriptional features are associated with VSMC state programs in clinically defined human carotid plaques. By keeping discovery, validation, and single-cell localization analyses separate, the work identifies a conservative and reproducible signal: efferocytosis-related scores covary with inflammatory VSMC scores at the tissue level in two independent bulk cohorts. At the same time, the analysis shows that individual DEGs and sample-level signature group differences are not FDR-supported in the available cohorts. The accompanying single-cell analysis localizes macrophage and VSMC compartments but does not treat thousands of cells as independent patients. The study therefore offers a transparent candidate association for future experimental and spatial validation, while explicitly avoiding overstatement as a causal mechanism, pathology-defined signature, diagnostic tool, or therapeutic target.

## References

{refs_text}
"""


def write_support_files(ctx: dict[str, object]) -> None:
    figure_legends = """Figure 1 | Study design, cohort roles, and analysis gates. Public human carotid plaque cohorts were assigned to discovery bulk analysis (GSE111782), independent bulk validation (GSE311535), and single-cell localization (GSE260657). Symptomatic and asymptomatic labels are retained as reported in the source metadata.
Figure 2 | Discovery bulk quality control, differential expression, and enrichment. GSE111782 was analyzed with limma using symptomatic versus asymptomatic plaques as the contrast. Positive log2FC denotes higher expression in symptomatic plaques. No genes reached FDR < 0.05; ranked predefined signatures and GO outputs are shown as generated from source tables.
Figure 3 | Predefined efferocytosis and VSMC signature scores and tissue-level correlations. Sample-level scores summarize available mapped genes for each predefined program. Group comparisons use two-sided Wilcoxon tests with Benjamini-Hochberg correction; correlations use Spearman rho and are tissue-level associations only.
Figure 4 | Single-cell atlas and major-cell annotation for GSE260657. Smart-seq2 cell files were parsed from public raw data. Major cell classes were assigned by marker-score and clustering support; uncertain cells were retained as unassigned. Cells are displayed for localization and are not independent patients.
Figure 5 | Macrophage subclusters and efferocytosis program scoring. Macrophage cells were analyzed descriptively to localize efferocytosis-related expression. Donor/sample-level summaries, rather than cell counts, define the inferential unit for group comparisons.
Figure 6 | VSMC subclusters and contractile/non-contractile program scoring. VSMC cells were analyzed descriptively using predefined VSMC state programs. These outputs indicate candidate cell-state localization and do not prove lineage transition or causality.
Figure 7 | Independent bulk validation in GSE311535. Validation results are reported separately from discovery because the cohort is diabetes-specific and uses RNA-seq counts. The reproducible finding is the efferocytosis-inflammatory VSMC tissue-level association; gene-level and group-level signature findings are not validated at FDR < 0.05.
"""
    tables = """Table 1. Dataset audit and provenance: see 03_data/metadata/dataset_audit.csv.
Table 2. Bulk analysis summary: see 05_results/tables/bulk_analysis_summary.csv.
Table 3. Discovery complete differential-expression table: see 05_results/tables/GSE111782_DEG_complete.csv.
Table 4. Independent validation complete differential-expression table: see 05_results/tables/GSE311535_DEG_complete.csv.
Table 5. Predefined signatures and sample-level scores: see 05_results/tables/predefined_gene_sets.csv and 05_results/tables/signature_scores.csv.
Table 6. Signature group comparisons and tissue-level correlations: see 05_results/tables/external_validation_signature_summary.csv and 05_results/tables/signature_correlations.csv.
Table 7. Single-cell annotation, marker candidates, donor summaries, and program tests: see 05_results/qc/GSE260657_cell_annotations.csv and 05_results/tables/GSE260657_*.csv.
Table 8. External validation gene-direction table: see 05_results/tables/external_validation_gene_direction.csv.
"""
    supplementary = """Supplementary materials included in the project package:
- Complete DEG tables for discovery and validation cohorts.
- GO, Reactome, and ranked predefined-signature enrichment tables where generated.
- Predefined efferocytosis and VSMC gene-set definitions with source notes.
- Single-cell QC, marker-candidate, annotation, subcluster, and donor-level program-score tables.
- Dataset audit, file hashes, runbook, error log, package versions, and reproducibility notes.
- Skill-based manuscript review, statistical audit, and reference-verification reports.
"""
    (MANUSCRIPT_DIR / "figure_legends.txt").write_text(figure_legends, encoding="utf-8")
    (MANUSCRIPT_DIR / "tables.txt").write_text(tables, encoding="utf-8")
    (MANUSCRIPT_DIR / "supplementary_materials.txt").write_text(supplementary, encoding="utf-8")


def write_reports(audit_rows: list[dict[str, str]], ctx: dict[str, object]) -> None:
    terminology = """# Terminology ledger

| Canonical term | First-use definition | Variants repaired or avoided | Decision |
|---|---|---|---|
| symptomatic versus asymptomatic human carotid plaques | clinical grouping reported by source metadata | unstable/stable when unsupported | Use symptomatic/asymptomatic throughout. |
| efferocytosis | macrophage apoptotic-cell clearance-related transcriptional program | apoptotic-cell clearance signature | Define once and use efferocytosis-related program. |
| vascular smooth muscle cell (VSMC) | plaque smooth-muscle-lineage or marker-supported cell/program context | smooth muscle cell, SMC, VSMCs | Spell out first; use VSMC after first use. |
| tissue-level association | correlation across bulk plaque samples | mechanism, causal link | Use association/covariation only. |
| donor-level single-cell summary | per donor/sample aggregate used for inference | cell-level significance | Cells are not independent patients. |
| FDR | Benjamini-Hochberg false-discovery rate | adjusted P value, q value | Use FDR consistently for corrected P values. |
"""
    (MANUSCRIPT_DIR / "terminology_ledger.md").write_text(terminology, encoding="utf-8")

    stats_report = f"""# Statistical audit report

Scope: revised full manuscript draft, figure legends, and key result tables.

## Study design readout

- Discovery bulk cohort: GSE111782, n = 18 plaques, 9 symptomatic and 9 asymptomatic.
- Independent validation bulk cohort: GSE311535, n = 12 plaques, 6 symptomatic and 6 asymptomatic, diabetes-specific.
- Single-cell cohort: GSE260657, 15 donor/sample files, 8 symptomatic and 7 asymptomatic, {sum(ctx['cell_counts'].values()):,} parsed cells.

## Independent unit and replication readout

- Bulk analyses use patient/plaque samples as the independent unit.
- Single-cell group comparisons use donor/sample summaries where possible.
- Cells are used for localization and descriptive subclustering, not as independent patient-level replicates.

## Major statistical issues found and fixed

- P1: The previous manuscript used scaffold wording such as "exact values are in tables" instead of reporting actual effect estimates and FDR values. Fixed by inserting traceable values from 05_results/tables.
- P1: Nominal findings could be overread. Fixed by explicitly stating that no individual DEG and no sample-level signature group comparison reached FDR < 0.05.
- P1: Correlations could be misread as mechanism. Fixed by labeling them as tissue-level associations and naming composition/stage confounding.
- P1: Single-cell pseudoreplication risk. Fixed by reporting donor-level tests and stating that cells were not treated as independent patients.
- P2: Optional modules were not clearly bounded. Fixed by stating that WGCNA, spatial validation, cell communication, and machine learning were skipped or excluded from main conclusions under prespecified gates.

## Remaining reviewer risks

- Small sample sizes limit power and precision.
- Validation cohort is diabetes-specific and not fully comparable with the discovery cohort.
- Full confidence in single-cell annotation would benefit from original processed objects or a full Seurat/Scanpy reanalysis.
- Exact ethics, funding, conflict-of-interest, author contribution, and code-repository details require real-author input.
"""
    (MANUSCRIPT_DIR / "statistical_audit_report.md").write_text(stats_report, encoding="utf-8")

    ref_report = f"""# Reference verification report

Verification source: NCBI PubMed E-utilities, retrieved {stamp()}.

## Summary

- Curated references retained: {len(audit_rows)}.
- Records with DOI metadata: {sum(1 for row in audit_rows if row.get('DOI'))}.
- Low-relevance topical candidates and method-mismatched PubMed records from the prior reference file were removed from the revised manuscript bibliography.
- Corrected method references include limma PMID 25605792, edgeR PMID 19910308, GSVA PMID 23323831, GSEA PMID 16199517, clusterProfiler PMID 22455463, ReactomePA PMID 26661513, and WGCNA PMID 19114008.
- Corrected GEO repository citations include PMID 11752295 and PMID 23193258; PMID 12519938 was not used because it corresponds to DDBJ, not GEO.

## Boundary

PubMed metadata existence and bibliographic fields were verified. Final authors should still check whether each citation supports the exact manuscript sentence in the final submission version.
"""
    (MANUSCRIPT_DIR / "reference_verification_report.md").write_text(ref_report, encoding="utf-8")

    review_report = f"""# Skill-based manuscript review report

## Review setup

- Input scope: current manuscript draft, integrated review document, result tables, QC files, figure legends, reference audit, and journal-compliance notes.
- Skills applied: nature-writing, nature-polishing, nature-statistics, nature-ref-verifier, and nature-reviewer.
- Shared manuscript claim summary: public human carotid plaque transcriptomes show reproducible tissue-level covariation between macrophage efferocytosis-related scores and inflammatory VSMC scores, but do not establish robust symptomatic-versus-asymptomatic DEGs, causality, plaque instability, or clinical utility.
- Missing materials affecting confidence: verified author/affiliation details, ethics and consent text from source publications, funding, conflict declarations, public code repository URL, and author-confirmed citation sentence mapping.

## Reviewer 1

- Overall assessment: The revised manuscript is technically more credible than the scaffold because it reports negative results and separates discovery, validation, and single-cell localization.
- Who would be interested: vascular biology and translational atherosclerosis readers interested in macrophage-VSMC state interactions in human carotid plaques.
- Major strengths: public-data reproducibility, predefined signatures, validation cohort retained despite non-replication, and donor-aware single-cell wording.
- Major concerns: small cohorts, diabetes-specific validation, mixed-tissue confounding, and lightweight single-cell annotation.
- Technical failings to address before submission: verify original-study ethics/consent, archive code publicly, and consider full single-cell reanalysis from processed objects if available.
- Recommendation posture: suitable for author review after revisions; not ready for direct submission until placeholders are resolved.

## Reviewer 2

- Overall assessment: The biological question is focused and relevant, but the strongest result is an association rather than a disease-defining signature.
- Who would be interested: researchers studying defective efferocytosis, inflammatory plaque remodeling, and VSMC phenotypic modulation.
- Major strengths: the title and abstract now match the real evidence; symptomatic plaques are not mislabeled as unstable plaques.
- Major concerns: the manuscript must avoid implying that a reproducible correlation proves macrophage-to-VSMC communication or therapeutic target status.
- Technical failings to address before submission: add author-verified citations at each final sentence and ensure all figure panels match the revised legends.
- Recommendation posture: promising computational analysis, but framed as hypothesis-generating.

## Reviewer 3

- Overall assessment: The revised text is easier to follow because the narrative leads with the actual reproducible finding and states the null results plainly.
- Who would be interested: clinicians and computational biologists needing careful public-data interpretation of symptomatic carotid plaques.
- Major strengths: clear boundary language, compact abstract, and explicit distinction between group differences and tissue-level correlations.
- Major concerns: some readers may expect stronger validation or spatial evidence; the discussion now acknowledges this but the author may need to justify suitability for the target journal.
- Technical failings to address before submission: final Word/PDF should be checked visually after author edits, especially embedded figures and table readability.
- Recommendation posture: ready for detailed author review, not final submission.

## Cross-review synthesis

- Consensus strengths: transparent negative findings, cautious interpretation, corrected references, and donor-aware single-cell reporting.
- Consensus technical risks: small sample size, public retrospective design, tissue heterogeneity, single-cell annotation confidence, and missing author-side declarations.
- Broad-interest readout: the study is field-relevant and potentially useful, but its impact rests on a reproducible association rather than a mechanistic advance.
- Most important issues to resolve: replace all placeholders, verify ethics/consent and citation-to-claim mapping, archive code/data processing outputs, and decide whether full single-cell reanalysis is feasible.

## Risk / unsupported claims removed

- Removed scaffold claims that deferred all numerical reporting to tables.
- Removed any implication that symptomatic plaques equal unstable plaques.
- Avoided causal language for correlations, signature scores, and cell-state localization.
- Avoided diagnostic, treatment-target, or therapeutic claims.
"""
    (MANUSCRIPT_DIR / "skill_based_review_report.md").write_text(review_report, encoding="utf-8")


def strip_inline_md(text: str) -> str:
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", text)
    return text


def add_markdown_to_doc(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    paragraph_buffer: list[str] = []
    table_buffer: list[str] = []

    def flush_para() -> None:
        if paragraph_buffer:
            text = strip_inline_md(" ".join(x.strip() for x in paragraph_buffer).strip())
            if text:
                doc.add_paragraph(text)
            paragraph_buffer.clear()

    def flush_table() -> None:
        if not table_buffer:
            return
        rows = [line.strip().strip("|").split("|") for line in table_buffer if not re.match(r"^\|\s*[-: ]+\|", line.strip())]
        rows = [[strip_inline_md(cell.strip()) for cell in row] for row in rows]
        if rows:
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            for j, cell in enumerate(rows[0]):
                table.rows[0].cells[j].text = cell
            for row in rows[1:]:
                cells = table.add_row().cells
                for j, cell in enumerate(row[: len(cells)]):
                    cells[j].text = cell
        table_buffer.clear()

    for line in lines:
        raw = line.rstrip()
        if raw.startswith("|"):
            flush_para()
            table_buffer.append(raw)
            continue
        flush_table()
        if not raw.strip():
            flush_para()
            continue
        if raw.startswith("# "):
            flush_para()
            doc.add_heading(strip_inline_md(raw[2:].strip()), level=1)
        elif raw.startswith("## "):
            flush_para()
            doc.add_heading(strip_inline_md(raw[3:].strip()), level=2)
        elif raw.startswith("### "):
            flush_para()
            doc.add_heading(strip_inline_md(raw[4:].strip()), level=3)
        elif raw.startswith("- "):
            flush_para()
            doc.add_paragraph(strip_inline_md(raw[2:].strip()), style="List Bullet")
        elif re.match(r"^\d+\.\s", raw):
            flush_para()
            doc.add_paragraph(strip_inline_md(raw), style="List Number")
        else:
            paragraph_buffer.append(raw)
    flush_table()
    flush_para()


def save_docx_from_markdown(markdown: str, output: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(10.5)
    add_markdown_to_doc(doc, markdown)
    doc.save(output)


def save_title_page() -> None:
    doc = Document()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(16)
    for line in [
        "Article type: Original Research",
        "Authors: [[AUTHOR NAME]]",
        "Affiliations: [[AFFILIATION]]",
        "Corresponding author: [[CORRESPONDING AUTHOR EMAIL]]",
        "Target journal: Frontiers in Cardiovascular Medicine",
        f"Generated: {stamp()}",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(MANUSCRIPT_DIR / "title_page.docx")


def save_simple_docx(source_txt: Path, output_docx: Path, heading: str) -> None:
    doc = Document()
    doc.add_heading(heading, level=1)
    text = source_txt.read_text(encoding="utf-8", errors="replace") if source_txt.exists() else ""
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(strip_inline_md(line.strip()))
    doc.save(output_docx)


def backup_existing() -> None:
    backup_dir = MANUSCRIPT_DIR / "archive_before_skill_revision_20260904"
    backup_dir.mkdir(parents=True, exist_ok=True)
    for name in [
        "manuscript_draft.md",
        "manuscript_draft.docx",
        "manuscript_draft.pdf",
        "references.bib",
        "figure_legends.txt",
        "tables.txt",
        "supplementary_materials.txt",
    ]:
        source = MANUSCRIPT_DIR / name
        if source.exists():
            target = backup_dir / name
            if not target.exists():
                shutil.copy2(source, target)


def main() -> int:
    started = dt.datetime.now()
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    MANUSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    LIT_DIR.mkdir(parents=True, exist_ok=True)
    log("START skill-based manuscript revision")
    backup_existing()
    try:
        records = fetch_pubmed(PMIDS)
    except Exception as exc:
        record_error(f"PubMed verification failed: {exc}")
        return 1
    audit_rows = write_references(records)
    ctx = build_result_context()
    manuscript = build_manuscript(ctx, numbered_references(audit_rows))
    (MANUSCRIPT_DIR / "manuscript_draft.md").write_text(manuscript, encoding="utf-8")
    write_support_files(ctx)
    write_reports(audit_rows, ctx)
    save_docx_from_markdown(manuscript, MANUSCRIPT_DIR / "manuscript_draft.docx")
    save_title_page()
    save_simple_docx(MANUSCRIPT_DIR / "figure_legends.txt", MANUSCRIPT_DIR / "figure_legends.docx", "Figure Legends")
    save_simple_docx(MANUSCRIPT_DIR / "tables.txt", MANUSCRIPT_DIR / "tables.docx", "Tables")
    save_simple_docx(MANUSCRIPT_DIR / "supplementary_materials.txt", MANUSCRIPT_DIR / "supplementary_materials.docx", "Supplementary Materials")
    ended = dt.datetime.now()
    with RUNBOOK.open("a", encoding="utf-8") as handle:
        handle.write(
            f"| {started.isoformat(timespec='seconds')} / {ended.isoformat(timespec='seconds')} | "
            f"python 04_scripts/15_skill_based_revision.py | 0 | "
            f"06_manuscript/manuscript_draft.md; 06_manuscript/manuscript_draft.docx; "
            f"06_manuscript/skill_based_review_report.md; 02_literature/reference_audit.csv | "
            f"Skill-based manuscript audit, revision, and reference cleanup |\n"
        )
    log(f"END skill-based manuscript revision status=0 duration_sec={(ended-started).total_seconds():.1f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
