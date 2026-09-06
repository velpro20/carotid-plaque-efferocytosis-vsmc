"""Generate a single Word review document for the carotid plaque manuscript.

Purpose: integrate title-page metadata, manuscript text, main figures, figure
legends, table/supplement inventories, key result summaries, and references.
Inputs: 06_manuscript/*.md/*.txt, 06_manuscript/references.bib,
05_results/figures/*.png, 05_results/tables/*.csv, 03_data/metadata/*.csv.
Outputs: 06_manuscript/integrated_review_document.docx and a run log.
Dependencies: python-docx, Pillow, Python standard library.
Run order: after 11_figure_generation.R and 12_manuscript_generation.R.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


SEED = 20260904
TITLE = (
    "Macrophage efferocytosis-related transcriptional features and VSMC state "
    "programs in symptomatic versus asymptomatic human carotid plaques"
)


def project_root() -> Path:
    if len(sys.argv) > 1:
        return Path(sys.argv[1]).resolve()
    return Path(__file__).resolve().parents[1]


ROOT = project_root()
MANUSCRIPT_DIR = ROOT / "06_manuscript"
FIGURE_DIR = ROOT / "05_results" / "figures"
TABLE_DIR = ROOT / "05_results" / "tables"
QC_DIR = ROOT / "05_results" / "qc"
LOG_DIR = ROOT / "08_logs"
METADATA_DIR = ROOT / "03_data" / "metadata"
OUTPUT = MANUSCRIPT_DIR / "integrated_review_document.docx"
LOG = LOG_DIR / "14_integrated_review_doc.log"
RUNBOOK = ROOT / "RUNBOOK.md"


def log(message: str) -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().isoformat(timespec="seconds")
    with LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"[{stamp}] {message}\n")


def read_text(path: Path) -> str:
    if not path.exists():
        log(f"MISSING {path}")
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        log(f"MISSING {path}")
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def add_heading(doc: Document, text: str, level: int = 1):
    para = doc.add_heading(text, level=level)
    return para


def add_note(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def strip_markdown_inline(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("__", "")
    text = text.replace("`", "")
    text = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1 (\2)", text)
    return text


def add_markdown_text(doc: Document, markdown: str, skip_title: bool = True) -> None:
    lines = markdown.splitlines()
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            text = strip_markdown_inline(" ".join(x.strip() for x in buffer).strip())
            if text:
                doc.add_paragraph(text)
            buffer.clear()

    for line in lines:
        raw = line.rstrip()
        if not raw.strip():
            flush()
            continue
        if skip_title and raw.startswith("# "):
            continue
        if raw.startswith("### "):
            flush()
            add_heading(doc, strip_markdown_inline(raw[4:].strip()), level=3)
        elif raw.startswith("## "):
            flush()
            add_heading(doc, strip_markdown_inline(raw[3:].strip()), level=2)
        elif raw.startswith("# "):
            flush()
            add_heading(doc, strip_markdown_inline(raw[2:].strip()), level=1)
        elif raw.startswith("- "):
            flush()
            doc.add_paragraph(strip_markdown_inline(raw[2:].strip()), style="List Bullet")
        else:
            buffer.append(raw)
    flush()


def add_small_table(doc: Document, rows: list[dict[str, str]], columns: list[str], title: str, max_rows: int | None = None) -> None:
    add_heading(doc, title, level=2)
    if not rows:
        doc.add_paragraph("No source rows were available for this table.")
        return
    rows_to_show = rows if max_rows is None else rows[:max_rows]
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, col in enumerate(columns):
        hdr[idx].text = col
    for row in rows_to_show:
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            cells[idx].text = str(row.get(col, ""))
    if max_rows is not None and len(rows) > max_rows:
        add_note(doc, f"Showing {max_rows} of {len(rows)} rows; complete table is available at {rows_source_hint(title)}.")


def rows_source_hint(title: str) -> str:
    mapping = {
        "Dataset Audit Summary": "03_data/metadata/dataset_audit.csv",
        "Bulk Analysis Summary": "05_results/tables/bulk_analysis_summary.csv",
        "Signature Group Comparisons": "05_results/tables/external_validation_signature_summary.csv",
        "Signature Correlations": "05_results/tables/signature_correlations.csv",
        "Single-Cell Donor-Level Program Tests": "05_results/tables/GSE260657_donor_program_tests.csv",
        "Single-Cell Cluster Annotation Summary": "05_results/qc/GSE260657_cluster_annotation_summary.csv",
    }
    return mapping.get(title, "the corresponding CSV file")


def fdr_count(path: Path) -> tuple[int, int]:
    rows = read_csv_rows(path)
    fdr_key = None
    if rows:
        for key in rows[0].keys():
            if key.lower() in {"fdr", "adj.p.val", "padj", "qvalue"}:
                fdr_key = key
                break
    if not fdr_key:
        return len(rows), 0
    sig = 0
    for row in rows:
        try:
            value = float(row.get(fdr_key, "nan"))
        except ValueError:
            continue
        if value < 0.05:
            sig += 1
    return len(rows), sig


def add_key_result_summary(doc: Document) -> None:
    add_heading(doc, "Computed Result Summary", level=1)
    add_note(
        doc,
        "This section is generated directly from project result tables for convenient review; "
        "complete source tables remain in 05_results/tables."
    )

    summary_rows = read_csv_rows(TABLE_DIR / "bulk_analysis_summary.csv")
    deg_rows: list[dict[str, str]] = []
    for cohort, filename in [
        ("GSE111782", "GSE111782_DEG_complete.csv"),
        ("GSE311535", "GSE311535_DEG_complete.csv"),
    ]:
        tested, sig = fdr_count(TABLE_DIR / filename)
        deg_rows.append({"cohort": cohort, "genes_tested_in_complete_table": str(tested), "FDR_lt_0.05_genes": str(sig)})
    add_small_table(doc, summary_rows, ["cohort", "model", "n_asymptomatic", "n_symptomatic", "n_genes_tested", "FDR_threshold"], "Bulk Analysis Summary")
    add_small_table(doc, deg_rows, ["cohort", "genes_tested_in_complete_table", "FDR_lt_0.05_genes"], "Bulk DEG FDR Summary")

    sig_group = read_csv_rows(TABLE_DIR / "external_validation_signature_summary.csv")
    add_small_table(
        doc,
        sig_group,
        ["cohort", "signature", "n", "median_asymptomatic", "median_symptomatic", "P_value", "FDR"],
        "Signature Group Comparisons",
    )

    corr_rows = read_csv_rows(TABLE_DIR / "signature_correlations.csv")
    eff_vsmc = [
        r for r in corr_rows
        if r.get("signature_1") == "efferocytosis" and str(r.get("signature_2", "")).endswith("vsmc")
    ]
    add_small_table(doc, eff_vsmc, ["cohort", "signature_1", "signature_2", "rho", "P_value", "FDR", "n_samples"], "Signature Correlations")

    sc_tests = read_csv_rows(TABLE_DIR / "GSE260657_donor_program_tests.csv")
    add_small_table(doc, sc_tests, ["cell_type", "program", "n_donors", "P_value", "FDR"], "Single-Cell Donor-Level Program Tests")

    sc_annot = read_csv_rows(QC_DIR / "GSE260657_cluster_annotation_summary.csv")
    add_small_table(
        doc,
        sc_annot,
        ["sc_cluster", "dominant_cell_type", "total_cells", "macrophage", "vsmc", "fibroblast", "endothelial", "t_nk", "unassigned"],
        "Single-Cell Cluster Annotation Summary",
    )


def add_figures(doc: Document) -> None:
    add_heading(doc, "Main Figures", level=1)
    figure_stems = [
        "Figure_1_study_design",
        "Figure_2_discovery_bulk",
        "Figure_3_signature_scores",
        "Figure_4_single_cell_atlas",
        "Figure_5_macrophage_efferocytosis",
        "Figure_6_vsmc_states",
        "Figure_7_external_validation",
    ]
    for stem in figure_stems:
        path = FIGURE_DIR / f"{stem}.png"
        add_heading(doc, stem.replace("_", " "), level=2)
        if path.exists():
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(str(path), width=Inches(6.5))
            add_note(doc, f"Source image: {path.relative_to(ROOT)}")
        else:
            doc.add_paragraph(f"Figure PNG was not found: {path.relative_to(ROOT)}")
            log(f"MISSING figure {path}")


def add_plain_text_section(doc: Document, heading: str, source: Path) -> None:
    add_heading(doc, heading, level=1)
    content = read_text(source)
    if not content:
        doc.add_paragraph(f"No content available from {source.relative_to(ROOT)}.")
        return
    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(strip_markdown_inline(line))
    add_note(doc, f"Source: {source.relative_to(ROOT)}")


def parse_bibtex(path: Path) -> list[dict[str, str]]:
    text = read_text(path)
    entries: list[dict[str, str]] = []
    for match in re.finditer(r"@\w+\s*\{([^,]+),(.*?)(?=\n\}\s*(?:\n@|\Z))", text, flags=re.S):
        body = match.group(2)
        entry = {"key": match.group(1).strip()}
        for field_match in re.finditer(r"\n\s*(\w+)\s*=\s*\{(.*?)\}\s*,?", body, flags=re.S):
            key = field_match.group(1).lower()
            value = re.sub(r"\s+", " ", field_match.group(2).strip())
            entry[key] = value
        if entry:
            entries.append(entry)
    return entries


def add_references(doc: Document) -> None:
    add_heading(doc, "References", level=1)
    refs = parse_bibtex(MANUSCRIPT_DIR / "references.bib")
    if not refs:
        doc.add_paragraph("No parsed references were available. See 06_manuscript/references.bib.")
        return
    add_note(
        doc,
        "Reference entries are expanded from references.bib for review. The final submission reference list "
        "should be reconciled against in-text citations and journal style."
    )
    for idx, ref in enumerate(refs, start=1):
        authors = ref.get("author", "")
        title = ref.get("title", "")
        journal = ref.get("journal", "")
        year = ref.get("year", "")
        doi = ref.get("doi", "")
        pmid = ref.get("pmid", "")
        parts = [f"{idx}. {authors}."]
        if title:
            parts.append(title)
        if journal or year:
            parts.append(f"{journal} {year}.".strip())
        if doi:
            parts.append(f"DOI: {doi}.")
        if pmid:
            parts.append(f"PMID: {pmid}.")
        doc.add_paragraph(" ".join(parts))


def add_file_inventory(doc: Document) -> None:
    add_heading(doc, "Source File Inventory", level=1)
    rows = []
    for path in sorted((MANUSCRIPT_DIR).glob("*")):
        if path.is_file() and path.name != OUTPUT.name:
            rows.append({"file": str(path.relative_to(ROOT)), "bytes": str(path.stat().st_size), "modified": dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")})
    for path in sorted(TABLE_DIR.glob("*.csv")):
        rows.append({"file": str(path.relative_to(ROOT)), "bytes": str(path.stat().st_size), "modified": dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")})
    add_small_table(doc, rows, ["file", "bytes", "modified"], "Review Source Files", max_rows=80)


def set_document_style(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def add_title_page(doc: Document) -> None:
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(16)

    for text in [
        "Target journal: Frontiers in Cardiovascular Medicine",
        "Article type: Original Research",
        "Integrated review document generated for author review",
        f"Generated on: {dt.datetime.now().isoformat(timespec='seconds')}",
        "Authors: [[AUTHOR NAME]]",
        "Affiliations: [[AFFILIATION]]",
        "Corresponding author: [[CORRESPONDING AUTHOR EMAIL]]",
        "Code repository: [[CODE REPOSITORY URL]]",
    ]:
        para = doc.add_paragraph(text)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    add_note(
        doc,
        "Review caution: this integrated file combines existing project outputs for easier reading. "
        "Author identity, ethics, funding, conflicts, repository URL, and final citation order require real-author verification before submission."
    )
    doc.add_page_break()


def main() -> int:
    started = dt.datetime.now()
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    MANUSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    log("START integrated review document generation")
    doc = Document()
    set_document_style(doc)
    add_title_page(doc)

    add_heading(doc, "Main Manuscript", level=1)
    add_markdown_text(doc, read_text(MANUSCRIPT_DIR / "manuscript_draft.md"), skip_title=True)
    doc.add_page_break()

    add_key_result_summary(doc)
    doc.add_page_break()

    for report_name, heading in [
        ("skill_based_review_report.md", "Skill-Based Review Report"),
        ("statistical_audit_report.md", "Statistical Audit Report"),
        ("reference_verification_report.md", "Reference Verification Report"),
        ("terminology_ledger.md", "Terminology Ledger"),
    ]:
        report_path = MANUSCRIPT_DIR / report_name
        if report_path.exists():
            add_plain_text_section(doc, heading, report_path)
            doc.add_page_break()

    add_figures(doc)
    doc.add_page_break()

    add_plain_text_section(doc, "Figure Legends", MANUSCRIPT_DIR / "figure_legends.txt")
    add_plain_text_section(doc, "Tables", MANUSCRIPT_DIR / "tables.txt")
    add_plain_text_section(doc, "Supplementary Materials", MANUSCRIPT_DIR / "supplementary_materials.txt")
    doc.add_page_break()

    add_small_table(
        doc,
        read_csv_rows(METADATA_DIR / "dataset_audit.csv"),
        ["accession", "dataset_title", "species", "tissue", "assay_type", "group_definition", "included_or_excluded", "intended_role", "limitations", "verification_status"],
        "Dataset Audit Summary",
        max_rows=30,
    )
    doc.add_page_break()

    add_file_inventory(doc)
    doc.add_page_break()

    add_references(doc)

    doc.save(OUTPUT)
    ended = dt.datetime.now()
    log(f"END integrated review document generation status=0 output={OUTPUT} duration_sec={(ended-started).total_seconds():.1f}")
    with RUNBOOK.open("a", encoding="utf-8") as handle:
        handle.write(
            f"| {started.isoformat(timespec='seconds')} / {ended.isoformat(timespec='seconds')} | "
            f"python 04_scripts/14_integrated_review_doc.py | 0 | "
            f"06_manuscript/integrated_review_document.docx; 08_logs/14_integrated_review_doc.log | "
            f"Generated consolidated Word review document |\n"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
